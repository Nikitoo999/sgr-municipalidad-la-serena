"""Base común para todos los ModelAdmin del proyecto SGR.
 
Todos los admins heredan de ``StandardAdmin`` para tener, de forma uniforme:
  * Paginación estándar (25 filas por página).
  * Filtro por fecha de creación en todas las listas.
  * Acción "Exportar selección a Excel (.xlsx)".
  * SweetAlert2 para validaciones, mensajes y confirmación de acciones masivas.
"""
from datetime import datetime
from decimal import Decimal
 
from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
 
SWEETALERT_CDN = "https://cdn.jsdelivr.net/npm/sweetalert2@11.10.8/dist/sweetalert2.all.min.js"
 
 
def _safe_text(value):
    """Evita inyección de fórmulas: un texto que empiece con = + - @ se
    interpretaría como fórmula al abrirlo en Excel."""
    if value and value[0] in "=+-@":
        return "'" + value
    return value
 
 
def _cell_value(obj, field):
    value = getattr(obj, field.name)
    if field.choices:
        return _safe_text(str(getattr(obj, f"get_{field.name}_display")()))
    if value is None:
        return ""
    if isinstance(value, bool):
        return "Sí" if value else "No"
    if isinstance(value, datetime):
        # Excel no admite fechas con zona horaria
        if timezone.is_aware(value):
            value = timezone.localtime(value).replace(tzinfo=None)
        return value
    if isinstance(value, Decimal):
        return float(value)
    if isinstance(value, (int, float)) or hasattr(value, "isoformat"):
        return value
    return _safe_text(str(value))
 
 
@admin.action(description="Exportar selección a Excel (.xlsx)")
def export_to_excel(modeladmin, request, queryset):
    opts = modeladmin.model._meta
    fields = list(opts.concrete_fields)
    fk_names = [f.name for f in fields if f.many_to_one]
    if fk_names:
        queryset = queryset.select_related(*fk_names)
 
    wb = Workbook()
    ws = wb.active
    ws.title = str(opts.verbose_name_plural)[:31]
 
    headers = []
    for f in fields:
        name = str(f.verbose_name)
        headers.append(name[:1].upper() + name[1:])
    ws.append(headers)
    for cell in ws[1]:
        cell.font = Font(bold=True)
    ws.freeze_panes = "A2"
 
    widths = [len(h) for h in headers]
    for obj in queryset.iterator():
        row = [_cell_value(obj, f) for f in fields]
        ws.append(row)
        for i, v in enumerate(row):
            widths[i] = max(widths[i], len(str(v)))
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = min(w + 2, 60)
 
    filename = f"{opts.model_name}_{timezone.localdate():%Y%m%d}.xlsx"
    response = HttpResponse(
        content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    response["Content-Disposition"] = f'attachment; filename="{filename}"'
    wb.save(response)
    return response
 
 
class StandardAdmin(admin.ModelAdmin):
    list_per_page = 25
    list_max_show_all = 100
    save_on_top = True
 
    class Media:
        js = (SWEETALERT_CDN, "sgr_core/admin_swal.js")
 
    def get_list_filter(self, request):
        filters = list(super().get_list_filter(request))
        if "created_at" not in filters:
            filters.append("created_at")
        return filters
 
    def get_actions(self, request):
        # Se agrega aquí (y no en `actions`) para que no se pierda cuando un
        # admin hijo define sus propias acciones.
        actions = super().get_actions(request)
        actions["export_to_excel"] = (
            export_to_excel,
            "export_to_excel",
            export_to_excel.short_description,
        )
        return actions