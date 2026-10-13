"""Pruebas de la Parte 6: mantenedor de Sub Atención.

Cubre: permisos por rol (solo superadmin y coordinador), código duplicado sin
distinguir mayúsculas, solo tipos activos en el formulario (conservando el
actual al editar aunque esté inactivo), regla de eliminación (con atenciones no
se elimina) y búsqueda/filtro/exportación.
"""

import io

from django.test import TestCase, override_settings
from django.urls import reverse

from openpyxl import load_workbook

from operaciones.models import Attention, Neighbor
from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
)
from sgr_core.models import AttentionSubtype, AttentionType


@override_settings(STORAGES=TEST_STORAGES)
class CrudAttentionSubtypesTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")

        cls.superuser = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
        )
        cls.delegado = make_employee(
            "delegado", "11.111.112-2", cls.d1, cls.roles["ROL03"],
        )
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
        )

        cls.tipo_activo = AttentionType.objects.create(code="TAT01", name="Social")
        cls.tipo_inactivo = AttentionType.objects.create(
            code="TAT05", name="Emergencias", status="Inactivo",
        )
        cls.sub_activa = AttentionSubtype.objects.create(
            code="SAT01", name="Ayuda social", attention_type=cls.tipo_activo,
        )

    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _flashed_text(self, resp):
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def _codes_shown(self, resp):
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _post(self, **extra):
        data = {
            "code": "SAT09", "name": "Nueva sub atención",
            "attention_type": str(self.tipo_activo.pk), "status": "Activo",
        }
        data.update(extra)
        return data

    def _crear_atencion(self, subtype):
        neighbor = Neighbor.objects.create(
            identifier="11.111.111-1", full_name="Vecino Demo",
        )
        return Attention.objects.create(
            neighbor=neighbor, delegation=self.d1,
            attention_type=self.tipo_activo, attention_subtype=subtype,
            description="Atención de prueba", responsible=self.superuser,
        )

    # ---------------------------------------------------------------
    # Permisos por rol
    # ---------------------------------------------------------------

    def test_solo_superadmin_y_coordinador(self):
        self._login("coordinador")
        self.assertEqual(self.client.get(reverse("attention_subtypes")).status_code, 200)
        for username in ("delegado", "funcionario"):
            self._login(username)
            self.assertEqual(
                self.client.get(reverse("attention_subtypes")).status_code, 403
            )

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("attention_subtypes"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    # ---------------------------------------------------------------
    # Código duplicado
    # ---------------------------------------------------------------

    def test_codigo_duplicado_sin_mayusculas(self):
        self._login("coordinador")
        resp = self.client.post(reverse("attention_subtypes_new"), self._post(code="sat01"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ya existe una sub atención con este código")

    # ---------------------------------------------------------------
    # Solo tipos activos en el formulario
    # ---------------------------------------------------------------

    def test_solo_tipos_activos_en_formulario(self):
        self._login("coordinador")
        resp = self.client.get(reverse("attention_subtypes_new"))
        form = resp.context["form"]
        codes = {t.code for t in form.fields["attention_type"].queryset}
        self.assertIn("TAT01", codes)
        self.assertNotIn("TAT05", codes)  # inactivo no se ofrece

    def test_editar_conserva_tipo_actual_inactivo(self):
        # Sub atención cuyo tipo quedó inactivo: al editar sigue ofreciéndose.
        sub = AttentionSubtype.objects.create(
            code="SAT09", name="Emergencia vial", attention_type=self.tipo_inactivo,
        )
        self._login("coordinador")
        resp = self.client.get(reverse("attention_subtypes_edit", args=[sub.pk]))
        form = resp.context["form"]
        codes = {t.code for t in form.fields["attention_type"].queryset}
        self.assertIn("TAT05", codes)  # el tipo actual inactivo se conserva

    # ---------------------------------------------------------------
    # Regla de eliminación
    # ---------------------------------------------------------------

    def test_sub_atencion_con_atenciones_no_se_elimina(self):
        sub = AttentionSubtype.objects.create(
            code="SAT02", name="Subsidio", attention_type=self.tipo_activo,
        )
        self._crear_atencion(sub)
        self._login("superadmin")
        resp = self.client.post(reverse("attention_subtypes_delete", args=[sub.pk]))
        self.assertEqual(resp.status_code, 302)
        sub.refresh_from_db()
        self.assertIsNone(sub.deleted_at)
        follow = self.client.get(reverse("attention_subtypes"))
        self.assertIn("Márcala como Inactiva", self._flashed_text(follow))

    def test_sub_atencion_sin_relaciones_se_elimina(self):
        sub = AttentionSubtype.objects.create(
            code="SAT03", name="Libre", attention_type=self.tipo_activo,
        )
        self._login("superadmin")
        resp = self.client.post(reverse("attention_subtypes_delete", args=[sub.pk]))
        self.assertEqual(resp.status_code, 302)
        sub.refresh_from_db()
        self.assertIsNotNone(sub.deleted_at)

    # ---------------------------------------------------------------
    # Búsqueda, filtro y exportación
    # ---------------------------------------------------------------

    def test_busqueda_filtro_exportacion(self):
        AttentionSubtype.objects.create(
            code="SAT08", name="Consulta general",
            attention_type=self.tipo_activo, status="Inactivo",
        )
        self._login("superadmin")
        resp = self.client.get(reverse("attention_subtypes"), {"q": "ayuda"})
        self.assertEqual(self._codes_shown(resp), ["SAT01"])

        resp = self.client.get(reverse("attention_subtypes"), {"estado": "Inactivo"})
        self.assertEqual(self._codes_shown(resp), ["SAT08"])

        resp = self.client.get(reverse("attention_subtypes_export"))
        sheet = load_workbook(filename=io.BytesIO(resp.content)).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["Código", "Nombre", "Tipo asociado", "Estado"])
        self.assertEqual(sheet.max_row - 1, AttentionSubtype.alive.count())

