"""Pruebas de la Parte 5, Etapa 2: mantenedor de Usuarios.

Cubre: alcance del listado por rango, crear (User + Employee en una
transaccion, username = correo), validacion de rango de rol en el servidor,
nadie cambia su propio rol, borrado logico que desactiva la cuenta, guard del
propio operador, exportacion con columnas configuradas y 403/404 con AuditLog.
"""

import io

from django.test import TestCase, override_settings
from django.urls import reverse

from openpyxl import load_workbook

from operaciones.models import Employee
from portal.models import AuditLog
from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
)


@override_settings(STORAGES=TEST_STORAGES)
class CrudUsersTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")
        cls.d2 = make_delegation("Avenida del Mar", "DEL02", color="#1D6FA5")

        cls.superuser = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
        )
        cls.coordinador.authorized_delegations.set([cls.d1, cls.d2])
        cls.delegado = make_employee(
            "delegado", "11.111.112-2", cls.d1, cls.roles["ROL03"],
        )
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
        )

    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _flashed_text(self, resp):
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def _create_post(self, role_code="ROL04", delegations=None, **extra):
        data = {
            "first_name": "Nora",
            "last_name": "Paredes",
            "email": "nora.paredes@ejemplo.cl",
            "rut": "12.345.678-5",
            "phone": "",
            "address": "",
            "role": str(self.roles[role_code].pk),
            "delegations": delegations if delegations is not None else [str(self.d1.pk)],
            "status": "Activo",
            "password1": "Demo2026!Segura",
            "password2": "Demo2026!Segura",
        }
        data.update(extra)
        return data

    def test_listado_acotado_por_rango(self):
        self._login("superadmin")
        resp = self.client.get(reverse("users"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(resp.context["total"], Employee.alive.count())

        self._login("delegado")
        resp = self.client.get(reverse("users"))
        self.assertEqual(resp.status_code, 200)
        for row in resp.context["rows"]:
            self.assertIn("ROL04", row["cells"][3]["text"])

        self._login("funcionario")
        self.assertEqual(self.client.get(reverse("users")).status_code, 403)

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("users"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    def test_403_url_directa_fuera_de_alcance(self):
        self._login("delegado")
        coord = Employee.alive.get(user__username="coordinador")
        self.assertEqual(
            self.client.get(reverse("users_detail", args=[coord.pk])).status_code, 403
        )
        denied = AuditLog.objects.filter(action="access_denied").count()
        self.assertGreaterEqual(denied, 1)

    def test_404_inexistente(self):
        self._login("superadmin")
        self.assertEqual(
            self.client.get(reverse("users_detail", args=[99999])).status_code, 404
        )

    def test_crear_consultar_editar(self):
        self._login("superadmin")
        resp = self.client.post(reverse("users_new"), self._create_post())
        self.assertEqual(resp.status_code, 302)
        created = Employee.alive.get(rut="12.345.678-5")
        self.assertEqual(created.user.username, "nora.paredes@ejemplo.cl")
        self.assertEqual(created.role.code, "ROL04")
        self.assertEqual(created.delegation, self.d1)
        self.assertTrue(created.user.is_active)
        self.assertTrue(created.user.check_password("Demo2026!Segura"))

        detail = self.client.get(reverse("users_detail", args=[created.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Consulta de registro")
        self.assertContains(detail, "disabled")

        resp = self.client.post(
            reverse("users_edit", args=[created.pk]),
            {
                "first_name": "Nora", "last_name": "Paredes Soto",
                "email": "nora.paredes@ejemplo.cl", "rut": "12.345.678-5",
                "phone": "", "address": "", "role": str(self.roles["ROL04"].pk),
                "delegations": [str(self.d1.pk)], "status": "Inactivo",
            },
        )
        self.assertEqual(resp.status_code, 302)
        created.refresh_from_db()
        self.assertEqual(created.user.last_name, "Paredes Soto")
        self.assertFalse(created.user.is_active)
        self.assertEqual(created.status_label, "Inactivo")

    def test_rol_fuera_de_rango_rechazado_en_servidor(self):
        self._login("delegado")
        resp = self.client.post(
            reverse("users_new"),
            self._create_post(role_code="ROL02", email="otro@ejemplo.cl",
                              rut="13.111.111-1"),
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "igual o superior", status_code=200)
        self.assertFalse(Employee.objects.filter(rut="13.111.111-1").exists())


    def test_nadie_cambia_su_propio_rol(self):
        self._login("superadmin")
        me = Employee.objects.create(
            user=self.superuser, rut="10.000.000-0",
            delegation=self.d1, role=self.roles["ROL04"],
        )
        resp = self.client.post(
            reverse("users_edit", args=[me.pk]),
            {
                "first_name": "", "last_name": "", "email": "superadmin@ejemplo.cl",
                "rut": "10.000.000-0", "phone": "", "address": "",
                "role": str(self.roles["ROL03"].pk),
                "delegations": [str(self.d1.pk)], "status": "Activo",
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "propio rol", status_code=200)

    def test_eliminar_desactiva_cuenta_y_no_es_fisico(self):
        self._login("superadmin")
        target = Employee.alive.get(user__username="funcionario")
        resp = self.client.post(reverse("users_delete", args=[target.pk]))
        self.assertEqual(resp.status_code, 302)
        target.refresh_from_db()
        self.assertIsNotNone(target.deleted_at)
        target.user.refresh_from_db()
        self.assertFalse(target.user.is_active)
        self.assertTrue(Employee.objects.filter(pk=target.pk).exists())

    def test_no_eliminar_propio_usuario(self):
        self._login("superadmin")
        me = Employee.objects.create(
            user=self.superuser, rut="10.000.001-1",
            delegation=self.d1, role=self.roles["ROL04"],
        )
        resp = self.client.post(reverse("users_delete", args=[me.pk]))
        self.assertEqual(resp.status_code, 302)
        me.refresh_from_db()
        self.assertIsNone(me.deleted_at)
        flashed = self._flashed_text(self.client.get(reverse("users")))
        self.assertIn("propio usuario", flashed)

    def test_eliminar_solo_post(self):
        self._login("superadmin")
        target = Employee.alive.get(user__username="funcionario")
        self.assertEqual(
            self.client.get(reverse("users_delete", args=[target.pk])).status_code, 404
        )

    def test_exportar_respeta_columnas_y_alcance(self):
        self._login("superadmin")
        resp = self.client.get(reverse("users_export"))
        self.assertEqual(resp.status_code, 200)
        sheet = load_workbook(filename=io.BytesIO(resp.content)).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(
            headers, ["RUT", "Nombre", "Correo", "Rol", "Delegaciones", "Estado"]
        )
        self.assertEqual(sheet.max_row - 1, Employee.alive.count())
        texts = [str(cell.value or "") for row in sheet.iter_rows(min_row=2) for cell in row]
        self.assertTrue(any("DEL01" in text for text in texts))

