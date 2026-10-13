"""Pruebas de la Parte 6: mantenedor de Tipo de Atención.

Cubre: permisos por rol (solo superadmin y coordinador), código duplicado sin
distinguir mayúsculas, reglas de eliminación (tipo con sub atenciones o con
atenciones no se elimina; sin relaciones se borra de forma lógica) y
búsqueda/filtro/exportación.
"""

import io

from django.test import TestCase, override_settings
from django.urls import reverse

from openpyxl import load_workbook

from operaciones.models import Attention, Neighbor
from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
)
from sgr_core.models import AttentionSubtype, AttentionType


@override_settings(STORAGES=TEST_STORAGES)
class CrudAttentionTypesTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")

        cls.superuser = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
        )
        cls.delegado = make_employee(
            "delegado", "11.111.112-2", cls.d1, cls.roles["ROL03"],
        )
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
        )

        cls.tipo_activo = AttentionType.objects.create(code="TAT01", name="Social")
        cls.tipo_inactivo = AttentionType.objects.create(
            code="TAT05", name="Emergencias", status="Inactivo",
        )
        # Sub atención vinculada al tipo activo (para la regla de eliminación).
        cls.sub = AttentionSubtype.objects.create(
            code="SAT01", name="Ayuda social", attention_type=cls.tipo_activo,
        )

    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _flashed_text(self, resp):
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def _codes_shown(self, resp):
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _post(self, **extra):
        data = {
            "code": "TAT06", "name": "Nuevo tipo",
            "description": "", "status": "Activo",
        }
        data.update(extra)
        return data

    def _crear_atencion(self, attention_type):
        neighbor = Neighbor.objects.create(
            identifier="11.111.111-1", full_name="Vecino Demo",
        )
        return Attention.objects.create(
            neighbor=neighbor, delegation=self.d1,
            attention_type=attention_type, attention_subtype=self.sub,
            description="Atención de prueba", responsible=self.superuser,
        )

    # ---------------------------------------------------------------
    # Permisos por rol
    # ---------------------------------------------------------------

    def test_solo_superadmin_y_coordinador(self):
        self._login("coordinador")
        self.assertEqual(self.client.get(reverse("attention_types")).status_code, 200)
        for username in ("delegado", "funcionario"):
            self._login(username)
            self.assertEqual(
                self.client.get(reverse("attention_types")).status_code, 403
            )

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("attention_types"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    # ---------------------------------------------------------------
    # Código duplicado
    # ---------------------------------------------------------------

    def test_codigo_duplicado_sin_mayusculas(self):
        self._login("coordinador")
        resp = self.client.post(reverse("attention_types_new"), self._post(code="tat01"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ya existe un tipo de atención con este código")

    # ---------------------------------------------------------------
    # Reglas de eliminación
    # ---------------------------------------------------------------

    def test_tipo_con_subatenciones_no_se_elimina(self):
        self._login("superadmin")
        resp = self.client.post(reverse("attention_types_delete", args=[self.tipo_activo.pk]))
        self.assertEqual(resp.status_code, 302)
        self.tipo_activo.refresh_from_db()
        self.assertIsNone(self.tipo_activo.deleted_at)
        follow = self.client.get(reverse("attention_types"))
        self.assertIn("Márcalo como Inactivo", self._flashed_text(follow))

    def test_tipo_con_atenciones_no_se_elimina(self):
        # TAT06 sin sub atenciones, pero con una atención asociada.
        tipo = AttentionType.objects.create(code="TAT06", name="Con atención")
        self._crear_atencion(tipo)
        self._login("superadmin")
        resp = self.client.post(reverse("attention_types_delete", args=[tipo.pk]))
        self.assertEqual(resp.status_code, 302)
        tipo.refresh_from_db()
        self.assertIsNone(tipo.deleted_at)
        follow = self.client.get(reverse("attention_types"))
        self.assertIn("atención(es)", self._flashed_text(follow))

    def test_tipo_sin_relaciones_se_elimina(self):
        tipo = AttentionType.objects.create(code="TAT07", name="Libre")
        self._login("superadmin")
        resp = self.client.post(reverse("attention_types_delete", args=[tipo.pk]))
        self.assertEqual(resp.status_code, 302)
        tipo.refresh_from_db()
        self.assertIsNotNone(tipo.deleted_at)

    # ---------------------------------------------------------------
    # Búsqueda, filtro y exportación
    # ---------------------------------------------------------------

    def test_busqueda_filtro_exportacion(self):
        self._login("superadmin")
        resp = self.client.get(reverse("attention_types"), {"q": "social"})
        self.assertEqual(self._codes_shown(resp), ["TAT01"])

        resp = self.client.get(reverse("attention_types"), {"estado": "Inactivo"})
        self.assertEqual(self._codes_shown(resp), ["TAT05"])

        resp = self.client.get(reverse("attention_types_export"))
        sheet = load_workbook(filename=io.BytesIO(resp.content)).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["Código", "Nombre", "Descripción", "Estado"])
        self.assertEqual(sheet.max_row - 1, AttentionType.alive.count())

