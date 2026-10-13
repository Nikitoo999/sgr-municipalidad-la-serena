"""Pruebas de la Parte 6: mantenedor de Ítems y metas.

Cubre: permisos por rol (coordinador gestiona, delegado solo ve/exporta,
funcionario sin acceso ni menú), código duplicado sin distinguir mayúsculas,
CRUD completo, búsqueda/filtro/exportación y borrado lógico.
"""

import io

from django.test import TestCase, override_settings
from django.urls import reverse

from openpyxl import load_workbook

from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
)
from sgr_core.models import ProjectItem


@override_settings(STORAGES=TEST_STORAGES)
class CrudItemsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")
        cls.d2 = make_delegation("Avenida del Mar", "DEL02")

        cls.superuser = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
        )
        cls.delegado = make_employee(
            "delegado", "11.111.112-2", cls.d1, cls.roles["ROL03"],
        )
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
        )

        cls.item_activo = ProjectItem.objects.create(
            code="ITM01", name="Operativos en terreno",
            description="Operativos municipales.",
        )
        cls.item_inactivo = ProjectItem.objects.create(
            code="ITM02", name="Solicitudes vecinales", status="Inactivo",
        )

    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _flashed_text(self, resp):
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def _codes_shown(self, resp):
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _post(self, **extra):
        data = {
            "code": "ITM09", "name": "Nuevo ítem",
            "description": "", "status": "Activo",
        }
        data.update(extra)
        return data

    # ---------------------------------------------------------------
    # Permisos por rol
    # ---------------------------------------------------------------

    def test_coordinador_y_delegado_ven_listado(self):
        self._login("coordinador")
        self.assertEqual(self.client.get(reverse("items")).status_code, 200)
        self._login("delegado")
        self.assertEqual(self.client.get(reverse("items")).status_code, 200)

    def test_delegado_solo_ve_y_exporta(self):
        self._login("delegado")
        self.assertEqual(self.client.get(reverse("items_new")).status_code, 403)
        self.assertEqual(
            self.client.get(reverse("items_edit", args=[self.item_activo.pk])).status_code,
            403,
        )
        self.assertEqual(self.client.get(reverse("items_export")).status_code, 200)

    def test_funcionario_sin_acceso_ni_menu(self):
        self._login("funcionario")
        self.assertEqual(self.client.get(reverse("items")).status_code, 403)
        home = self.client.get(reverse("home"))
        labels = [item["label"] for item in home.context["sgr_menu"]]
        self.assertNotIn("Ítems y metas", labels)
        self.assertNotIn("Tipo de Atención", labels)
        self.assertNotIn("Sub Atención", labels)

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("items"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    # ---------------------------------------------------------------
    # Código duplicado y CRUD
    # ---------------------------------------------------------------

    def test_codigo_duplicado_sin_mayusculas(self):
        self._login("coordinador")
        resp = self.client.post(reverse("items_new"), self._post(code="itm01"))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ya existe un ítem con este código", status_code=200)

    def test_crear_consultar_editar(self):
        self._login("coordinador")
        resp = self.client.post(reverse("items_new"), self._post())
        self.assertEqual(resp.status_code, 302)
        created = ProjectItem.alive.get(code="ITM09")
        self.assertEqual(created.name, "Nuevo ítem")

        detail = self.client.get(reverse("items_detail", args=[created.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Consulta de registro")

        resp = self.client.post(
            reverse("items_edit", args=[created.pk]),
            self._post(name="Ítem editado"),
        )
        self.assertEqual(resp.status_code, 302)
        created.refresh_from_db()
        self.assertEqual(created.name, "Ítem editado")

    def test_eliminar_logico(self):
        self._login("superadmin")
        resp = self.client.post(reverse("items_delete", args=[self.item_inactivo.pk]))
        self.assertEqual(resp.status_code, 302)
        self.item_inactivo.refresh_from_db()
        self.assertIsNotNone(self.item_inactivo.deleted_at)

    # ---------------------------------------------------------------
    # Búsqueda, filtro y exportación
    # ---------------------------------------------------------------

    def test_busqueda_filtro_exportacion(self):
        self._login("superadmin")
        resp = self.client.get(reverse("items"), {"q": "operativo"})
        self.assertEqual(self._codes_shown(resp), ["ITM01"])

        resp = self.client.get(reverse("items"), {"estado": "Inactivo"})
        self.assertEqual(self._codes_shown(resp), ["ITM02"])

        resp = self.client.get(reverse("items_export"))
        sheet = load_workbook(filename=io.BytesIO(resp.content)).active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["Código", "Nombre", "Descripción", "Estado"])
        self.assertEqual(sheet.max_row - 1, ProjectItem.alive.count())

