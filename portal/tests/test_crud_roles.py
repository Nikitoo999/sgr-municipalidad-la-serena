"""Pruebas de la Parte 5, Etapa 1: mantenedor de Roles.

Cubre: permisos por rol (solo el superadmin entra), matriz guardada con efecto
inmediato en el acceso, "Ver" obligatorio, sin módulos habilitados, rol del
sistema no editable ni eliminable, rol con usuarios no eliminable ni
desactivable, código duplicado sin distinguir mayúsculas, 403/404 con AuditLog,
exportación solo con las columnas configuradas y la no-regresión del ancho
completo de los campos de Delegaciones (formulario base del motor).
"""

import io

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse

from openpyxl import load_workbook

from portal.models import AuditLog
from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
    permission_pairs,
)
from sgr_core.models import Role, RolePermission


@override_settings(STORAGES=TEST_STORAGES)
class CrudRolesTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")
        cls.d2 = make_delegation("Avenida del Mar", "DEL02", color="#1D6FA5")

        cls.superuser = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
        )
        cls.delegado = make_employee(
            "delegado", "11.111.112-2", cls.d1, cls.roles["ROL03"],
        )
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
        )
        cls._rut_seq = 113

    # ---------------------------------------------------------------
    # Utilidades
    # ---------------------------------------------------------------

    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _flashed_text(self, resp):
        """Texto de los messages de Django (solo SweetAlert2 los muestra)."""
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def _codes_shown(self, resp):
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _matrix_post(self, **extra):
        """POST base válido: un rol con Usuarios: Ver, crear, editar, exportar."""
        data = {
            "code": "ROL05",
            "name": "Rol de prueba",
            "description": "",
            "scope": "own_delegation",
            "status": "Activo",
            "permissions": [
                "users.view", "users.create", "users.edit", "users.export",
            ],
        }
        data.update(extra)
        return data

    # ---------------------------------------------------------------
    # Permisos por rol
    # ---------------------------------------------------------------

    def test_listado_solo_superadmin(self):
        self._login("superadmin")
        self.assertEqual(self.client.get(reverse("roles")).status_code, 200)

        for username in ("coordinador", "delegado", "funcionario"):
            self._login(username)
            resp = self.client.get(reverse("roles"))
            self.assertEqual(resp.status_code, 403)
            self.assertContains(resp, "Acceso restringido", status_code=403)

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("roles"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    def test_crear_consultar_editar_rol(self):
        self._login("superadmin")
        resp = self.client.post(reverse("roles_new"), self._matrix_post())
        self.assertEqual(resp.status_code, 302)
        created = Role.alive.get(code="ROL05")
        self.assertEqual(created.scope, "own_delegation")
        self.assertEqual(
            permission_pairs(created),
            {("users", "view"), ("users", "create"), ("users", "edit"), ("users", "export")},
        )

        detail = self.client.get(reverse("roles_detail", args=[created.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Consulta de registro")
        # Modo consulta: la matriz se renderiza deshabilitada.
        self.assertContains(detail, "disabled")

        resp = self.client.post(
            reverse("roles_edit", args=[created.pk]),
            self._matrix_post(name="Rol editado", permissions=["items.view"]),
        )
        self.assertEqual(resp.status_code, 302)
        created.refresh_from_db()
        self.assertEqual(created.name, "Rol editado")
        self.assertEqual(permission_pairs(created), {("items", "view")})

    def test_matriz_guarda_permisos_y_rige_de_inmediato(self):
        """La matriz editada se lee de la BD: rige desde la siguiente solicitud."""
        self._login("superadmin")
        rol = Role.alive.get(code="ROL03")
        # El delegado hoy ve Usuarios.
        self._login("delegado")
        self.assertEqual(self.client.get(reverse("users")).status_code, 200)

        # Se le quita el permiso 'view' de usuarios al rol del delegado.
        self._login("superadmin")
        resp = self.client.post(
            reverse("roles_edit", args=[rol.pk]),
            {
                "code": rol.code,
                "name": rol.name,
                "description": rol.description,
                "scope": rol.scope,
                "status": "Activo",
                "permissions": ["attentions.view", "attentions.create"],
            },
        )
        self.assertEqual(resp.status_code, 302)
        self.assertEqual(
            permission_pairs(rol), {("attentions", "view"), ("attentions", "create")},
        )

        # Al reingresar como delegado, el menú ya no ofrece Usuarios (403).
        self._login("delegado")
        home = self.client.get(reverse("home"))
        self.assertEqual(home.status_code, 200)
        self.assertNotIn("Usuarios", [item["label"] for item in home.context["sgr_menu"]])
        self.assertEqual(self.client.get(reverse("users")).status_code, 403)

    def test_modulo_sin_ver_rechazado(self):
        self._login("superadmin")
        resp = self.client.post(
            reverse("roles_new"),
            self._matrix_post(permissions=["users.create", "users.edit"]),
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "debe incluir", status_code=200)
        self.assertFalse(Role.alive.filter(code="ROL05").exists())

    def test_sin_modulos_habilitados_rechazado(self):
        self._login("superadmin")
        resp = self.client.post(reverse("roles_new"), self._matrix_post(permissions=[]))
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Habilita al menos un módulo", status_code=200)
        self.assertFalse(Role.alive.filter(code="ROL05").exists())

    def test_codigo_duplicado_sin_distinguir_mayusculas(self):
        self._login("superadmin")
        resp = self.client.post(
            reverse("roles_new"), self._matrix_post(code="rol01")
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ya existe un rol con este código", status_code=200)
        # Se guarda en mayúsculas y no duplica el ROL01 existente.
        self.assertEqual(Role.objects.filter(code__iexact="ROL01").count(), 1)

    # ---------------------------------------------------------------
    # Rol del sistema: solo consulta
    # ---------------------------------------------------------------

    def test_rol_sistema_no_editable(self):
        self._login("superadmin")
        rol = Role.alive.get(code="ROL01")
        self.assertTrue(rol.is_system)

        resp = self.client.post(
            reverse("roles_edit", args=[rol.pk]),
            self._matrix_post(code="ROL01", name="Superadministrador"),
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "solo se puede consultar", status_code=200)
        self.assertContains(resp, "perfil base del sistema", status_code=200)
        rol.refresh_from_db()
        self.assertEqual(rol.name, "Superadministrador")
        # La matriz original del sistema queda intacta (7 módulos × 5 acciones).
        self.assertEqual(len(permission_pairs(rol)), 35)

    def test_rol_sistema_no_eliminable(self):
        self._login("superadmin")
        rol = Role.alive.get(code="ROL01")
        resp = self.client.post(reverse("roles_delete", args=[rol.pk]))
        self.assertEqual(resp.status_code, 302)
        rol.refresh_from_db()
        self.assertIsNone(rol.deleted_at)
        flashed = self._flashed_text(self.client.get(reverse("roles")))
        self.assertIn("perfil base del sistema", flashed)

    # ---------------------------------------------------------------
    # Rol con usuarios: no se elimina ni se desactiva
    # ---------------------------------------------------------------

    def test_rol_con_usuarios_no_se_elimina(self):
        self._login("superadmin")
        rol = Role.alive.get(code="ROL03")
        self.assertEqual(rol.employees.count(), 1)

        resp = self.client.post(reverse("roles_delete", args=[rol.pk]))
        self.assertEqual(resp.status_code, 302)
        rol.refresh_from_db()
        self.assertIsNone(rol.deleted_at)
        flashed = self._flashed_text(self.client.get(reverse("roles")))
        self.assertIn("1 usuario(s) asignado(s)", flashed)
        self.assertIn("Reasínalos", flashed)

    def test_rol_con_usuarios_no_se_desactiva(self):
        self._login("superadmin")
        rol = Role.alive.get(code="ROL03")
        resp = self.client.post(
            reverse("roles_edit", args=[rol.pk]),
            {
                "code": rol.code,
                "name": rol.name,
                "description": rol.description,
                "scope": rol.scope,
                "status": "Inactivo",
                "permissions": ["users.view"],
            },
        )
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "usuario(s) asignado(s)", status_code=200)
        self.assertContains(resp, "Reasínalos", status_code=200)
        rol.refresh_from_db()
        self.assertEqual(rol.status, "Activo")

    def test_rol_sin_usuarios_se_puede_desactivar(self):
        self._login("superadmin")
        self.client.post(reverse("roles_new"), self._matrix_post())
        created = Role.alive.get(code="ROL05")
        resp = self.client.post(
            reverse("roles_edit", args=[created.pk]),
            self._matrix_post(status="Inactivo"),
        )
        self.assertEqual(resp.status_code, 302)
        created.refresh_from_db()
        self.assertEqual(created.status, "Inactivo")

    # ---------------------------------------------------------------
    # 403 / 404 con auditoría
    # ---------------------------------------------------------------

    def test_403_url_directa_otro_rol_con_auditlog(self):
        self._login("delegado")
        rol = Role.alive.get(code="ROL03")
        self.assertEqual(
            self.client.get(reverse("roles_detail", args=[rol.pk])).status_code, 403
        )
        self.assertEqual(
            self.client.get(reverse("roles_edit", args=[rol.pk])).status_code, 403
        )
        self.assertTrue(
            AuditLog.objects.filter(action="access_denied", module="roles").exists()
        )

    def test_404_inexistente(self):
        self._login("superadmin")
        self.assertEqual(
            self.client.get(reverse("roles_detail", args=[99999])).status_code, 404
        )
        self.assertEqual(
            self.client.get(reverse("roles_edit", args=[99999])).status_code, 404
        )

    # ---------------------------------------------------------------
    # Exportación
    # ---------------------------------------------------------------

    def test_exportacion_columnas_configuradas(self):
        self._login("superadmin")
        resp = self.client.get(reverse("roles_export"))
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("roles_", resp["Content-Disposition"])

        workbook = load_workbook(io.BytesIO(resp.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(
            headers, ["Código", "Rol", "Alcance", "Módulos habilitados", "Estado"]
        )
        # Nunca columnas no configuradas (la descripción no se exporta).
        self.assertNotIn("Descripción", headers)
        # Exporta todos los roles vivos, no solo la página visible.
        self.assertEqual(sheet.max_row - 1, Role.alive.count())

    # ---------------------------------------------------------------
    # No-regresión del formulario base del motor (Delegaciones)
    # ---------------------------------------------------------------

    def test_delegaciones_conserva_campos_de_ancho_completo(self):
        """El refactor a ``full_width_fields`` no cambia el ancho de Delegaciones."""
        self._login("superadmin")
        resp = self.client.get(reverse("delegations_new"))
        self.assertEqual(resp.status_code, 200)

        # El formulario base conserva el valor por defecto de la Parte 4.
        self.assertEqual(
            tuple(resp.context["form"].full_width_fields), ("description", "address"),
        )
        # DelegationForm no declara 'description': su único campo de ancho
        # completo sigue siendo 'address', igual que antes del refactor.
        self.assertNotIn("description", resp.context["form"].fields)

        import re

        full_fields = re.findall(
            r'<div class="full">\s*<label class="fl" for="id_(\w+)"',
            resp.content.decode(),
        )
        self.assertEqual(full_fields, ["address"])
        # El resto de los campos (código incluido) sigue en media columna.
        self.assertIn('<div class="">', resp.content.decode())



