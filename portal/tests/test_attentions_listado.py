"""Pruebas del listado avanzado de Atenciones (filtros, orden y page_size).

Cubre el listado (``AttentionListView``) y la exportación a Excel
(``AttentionExportView``): ambos leen los mismos parámetros GET. Verifica que
los valores inválidos se ignoran sin error, que el alcance por rol no se puede
ampliar con parámetros y que la exportación coincide con el listado.
"""

import io
import re
from datetime import datetime

from django.db import connection
from django.test import TestCase, override_settings
from django.test.utils import CaptureQueriesContext
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from operaciones.models import Attention, Neighbor
from portal.tests.helpers import (
    DEMO_PASSWORD,
    TEST_STORAGES,
    ensure_roles,
    make_delegation,
    make_employee,
    make_user,
)
from sgr_core.models import AttentionSubtype, AttentionType


def aware(year, month, day, hour=12):
    return timezone.make_aware(datetime(year, month, day, hour, 0))


@override_settings(STORAGES=TEST_STORAGES)
class AttentionsListadoTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        cls.roles = ensure_roles()
        cls.d1 = make_delegation("Centro", "DEL01", color="#C62828")
        cls.d2 = make_delegation("Avenida del Mar", "DEL02", color="#1D6FA5")

        cls.superadmin = make_user(
            "superadmin", password=DEMO_PASSWORD, is_superuser=True, is_staff=True,
        )
        cls.coordinador = make_employee(
            "coordinador", "11.111.111-1", cls.d1, cls.roles["ROL02"],
            first_name="Carla", last_name="Coordina",
        )
        cls.coordinador.authorized_delegations.set([cls.d1, cls.d2])
        cls.funcionario = make_employee(
            "funcionario", "11.111.113-3", cls.d1, cls.roles["ROL04"],
            first_name="Felipe", last_name="Funciona",
        )
        cls.otro = make_employee(
            "otro.funcionario", "11.111.114-4", cls.d1, cls.roles["ROL04"],
            first_name="Ana", last_name="Alvarez",
        )

        cls.tipo_a = AttentionType.objects.create(code="TAT01", name="Social")
        cls.tipo_b = AttentionType.objects.create(code="TAT02", name="Territorial")
        cls.sub_a1 = AttentionSubtype.objects.create(
            code="SAT01", name="Ayuda social", attention_type=cls.tipo_a)
        cls.sub_a2 = AttentionSubtype.objects.create(
            code="SAT02", name="Visita domiciliaria", attention_type=cls.tipo_a)
        cls.sub_b1 = AttentionSubtype.objects.create(
            code="SAT03", name="Operativo", attention_type=cls.tipo_b)

        cls.maria = Neighbor.objects.create(identifier="11.111.111-1", full_name="María González")
        cls.pedro = Neighbor.objects.create(identifier="22.222.222-2", full_name="Pedro Soto")

        def make(neighbor, delegation, tipo, sub, responsible, status, day):
            att = Attention.objects.create(
                neighbor=neighbor, delegation=delegation,
                attention_type=tipo, attention_subtype=sub,
                description="Descripción de prueba de la atención.",
                responsible=responsible.user if hasattr(responsible, "user") else responsible,
                status=status,
            )
            Attention.objects.filter(pk=att.pk).update(created_at=aware(2026, 3, day))
            att.refresh_from_db()
            return att

        S = Attention.AttentionStatus
        # N° automático ATN-0001.. en orden de creación
        cls.a1 = make(cls.maria, cls.d1, cls.tipo_a, cls.sub_a1, cls.funcionario, S.RESOLVED, 10)
        cls.a2 = make(cls.pedro, cls.d1, cls.tipo_a, cls.sub_a2, cls.otro, S.RECEIVED, 20)
        cls.a3 = make(cls.maria, cls.d1, cls.tipo_b, cls.sub_b1, cls.funcionario, S.IN_PROGRESS, 5)
        cls.a4 = make(cls.pedro, cls.d2, cls.tipo_b, cls.sub_b1, cls.otro, S.RECEIVED, 15)

    # --- utilidades -------------------------------------------------------
    def _login(self, username):
        self.client.logout()
        assert self.client.login(username=username, password=DEMO_PASSWORD)

    def _get(self, **params):
        return self.client.get(reverse("attentions"), params)

    def _numbers(self, resp):
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _export_numbers(self, **params):
        resp = self.client.get(reverse("attentions_export"), params)
        self.assertEqual(resp.status_code, 200)
        sheet = load_workbook(filename=io.BytesIO(resp.content)).active
        return [row[0].value for row in sheet.iter_rows(min_row=2)]

    def _bulk(self, count):
        """Crea ``count`` atenciones extra (delegación Centro, tipo A)."""
        for _ in range(count):
            Attention.objects.create(
                neighbor=self.maria, delegation=self.d1,
                attention_type=self.tipo_a, attention_subtype=self.sub_a1,
                description="Descripción de prueba de la atención.",
                responsible=self.funcionario.user,
            )

    # --- orden ------------------------------------------------------------
    def test_orden_por_defecto_es_numero_ascendente(self):
        self._login("superadmin")
        self.assertEqual(self._numbers(self._get()),
                         [self.a1.number, self.a2.number, self.a3.number, self.a4.number])

    def test_orden_numero_descendente(self):
        self._login("superadmin")
        self.assertEqual(self._numbers(self._get(orden="-numero")),
                         [self.a4.number, self.a3.number, self.a2.number, self.a1.number])

    def test_orden_por_fecha(self):
        self._login("superadmin")
        # fechas: a3 (5), a1 (10), a4 (15), a2 (20)
        self.assertEqual(self._numbers(self._get(orden="fecha")),
                         [self.a3.number, self.a1.number, self.a4.number, self.a2.number])
        self.assertEqual(self._numbers(self._get(orden="-fecha")),
                         [self.a2.number, self.a4.number, self.a1.number, self.a3.number])

    def test_orden_por_solicitante(self):
        self._login("superadmin")
        numbers = self._numbers(self._get(orden="solicitante"))
        # María (a1, a3) antes que Pedro (a2, a4); desempate estable por pk
        self.assertEqual(numbers, [self.a1.number, self.a3.number, self.a2.number, self.a4.number])

    def test_orden_por_tipo_subtipo_delegacion_responsable(self):
        self._login("superadmin")
        for key in ("tipo", "subtipo", "delegacion", "responsable"):
            asc = self._numbers(self._get(orden=key))
            desc = self._numbers(self._get(orden=f"-{key}"))
            self.assertEqual(len(asc), 4, key)
            self.assertEqual(sorted(asc), sorted(desc), key)
            self.assertNotEqual(asc, desc, key)
        # delegación: DEL01 (a1,a2,a3) antes que DEL02 (a4)
        self.assertEqual(self._numbers(self._get(orden="delegacion"))[-1], self.a4.number)
        # responsable: Ana Alvarez (otro: a2, a4) antes que Felipe (a1, a3)
        self.assertEqual(self._numbers(self._get(orden="responsable")),
                         [self.a2.number, self.a4.number, self.a1.number, self.a3.number])

    def test_orden_por_estado_sigue_el_avance(self):
        self._login("superadmin")
        # Ingresada (a2, a4) -> En gestión (a3) -> Resuelta (a1)
        self.assertEqual(self._numbers(self._get(orden="estado")),
                         [self.a2.number, self.a4.number, self.a3.number, self.a1.number])
        self.assertEqual(self._numbers(self._get(orden="-estado"))[0], self.a1.number)

    def test_orden_invalido_se_ignora(self):
        self._login("superadmin")
        for bad in ("zzz", "-", "numero;DROP TABLE atencion", "--numero", "number", "  "):
            resp = self._get(orden=bad)
            self.assertEqual(resp.status_code, 200, bad)
            self.assertEqual(self._numbers(resp)[0], self.a1.number, bad)

    def test_encabezados_ordenables_con_aria_sort(self):
        self._login("superadmin")
        html = self._get(orden="-fecha").content.decode()
        self.assertIn('aria-sort="descending"', html)
        self.assertIn('aria-sort="none"', html)
        # un clic en la columna actual (desc) vuelve a ascendente
        resp = self._get(orden="-fecha")
        fecha = next(h for h in resp.context["headers"] if h["label"] == "Fecha y hora")
        self.assertIn("orden=fecha", fecha["url"])
        self.assertNotIn("orden=-fecha", fecha["url"])
        # y la columna ascendente pasa a descendente
        resp = self._get(orden="fecha")
        fecha = next(h for h in resp.context["headers"] if h["label"] == "Fecha y hora")
        self.assertIn("orden=-fecha", fecha["url"])

    def test_cambiar_orden_vuelve_a_pagina_1(self):
        self._login("superadmin")
        self._bulk(8)
        resp = self._get(page=2, orden="-numero")
        header = next(h for h in resp.context["headers"] if h["label"] == "Tipo")
        self.assertNotIn("page=", header["url"])

    # --- filtros ----------------------------------------------------------
    def test_filtro_por_tipo(self):
        self._login("superadmin")
        self.assertEqual(sorted(self._numbers(self._get(tipo=self.tipo_a.pk))),
                         sorted([self.a1.number, self.a2.number]))

    def test_filtro_por_subtipo(self):
        self._login("superadmin")
        self.assertEqual(sorted(self._numbers(self._get(subtipo=self.sub_b1.pk))),
                         sorted([self.a3.number, self.a4.number]))

    def test_filtro_tipo_y_subtipo_incompatibles_da_vacio(self):
        self._login("superadmin")
        resp = self._get(tipo=self.tipo_a.pk, subtipo=self.sub_b1.pk)
        self.assertEqual(self._numbers(resp), [])
        self.assertContains(resp, "No encontramos resultados")

    def test_opciones_de_subtipo_dependen_del_tipo(self):
        self._login("superadmin")
        resp = self._get(tipo=self.tipo_a.pk)
        labels = [o["label"] for o in resp.context["subtype_options"]]
        self.assertEqual(labels, ["Ayuda social", "Visita domiciliaria"])

    def test_filtro_por_responsable(self):
        self._login("superadmin")
        resp = self._get(responsable=self.otro.user.pk)
        self.assertEqual(sorted(self._numbers(resp)), sorted([self.a2.number, self.a4.number]))

    def test_filtro_por_rango_de_fechas(self):
        self._login("superadmin")
        # a1 (10 mar) y a4 (15 mar)
        resp = self._get(desde="2026-03-08", hasta="2026-03-16")
        self.assertEqual(sorted(self._numbers(resp)), sorted([self.a1.number, self.a4.number]))

    def test_rango_de_fechas_incluye_el_dia_hasta_completo(self):
        self._login("superadmin")
        resp = self._get(desde="2026-03-20", hasta="2026-03-20")   # a2 a las 12:00
        self.assertEqual(self._numbers(resp), [self.a2.number])

    def test_solo_desde_y_solo_hasta(self):
        self._login("superadmin")
        self.assertEqual(sorted(self._numbers(self._get(desde="2026-03-15"))),
                         sorted([self.a4.number, self.a2.number]))
        self.assertEqual(sorted(self._numbers(self._get(hasta="2026-03-10"))),
                         sorted([self.a3.number, self.a1.number]))

    def test_rango_invertido_avisa_y_no_filtra(self):
        self._login("superadmin")
        resp = self._get(desde="2026-03-20", hasta="2026-03-01")
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(self._numbers(resp)), 4)
        self.assertContains(resp, "no puede ser posterior")

    def test_parametros_invalidos_se_ignoran_sin_error(self):
        self._login("superadmin")
        resp = self._get(
            tipo="abc", subtipo="-3", delegacion="x", responsable="1;DROP",
            desde="31/12/2026", hasta="no-es-fecha", estado="inventado",
            por_pagina="999", page="zzz",
        )
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(len(self._numbers(resp)), 4)

    def test_filtros_combinados_con_busqueda(self):
        self._login("superadmin")
        resp = self._get(q="maría", tipo=self.tipo_b.pk, estado="in_progress")
        self.assertEqual(self._numbers(resp), [self.a3.number])

    def test_filtro_de_delegacion_solo_en_vista_todas(self):
        self._login("superadmin")      # delegación activa por defecto: "Todas"
        resp = self._get(delegacion=self.d2.pk)
        self.assertTrue(resp.context["show_delegation_filter"])
        self.assertEqual(self._numbers(resp), [self.a4.number])

    def test_delegacion_ignorada_con_delegacion_activa_fija(self):
        # El coordinador tiene activa DEL01: el parámetro no amplía ni vacía.
        self._login("coordinador")
        resp = self._get(delegacion=self.d2.pk)
        self.assertFalse(resp.context["show_delegation_filter"])
        self.assertEqual(sorted(self._numbers(resp)),
                         sorted([self.a1.number, self.a2.number, self.a3.number]))

    # --- alcance por rol no ampliable ------------------------------------
    def test_funcionario_no_ve_ajenas_aunque_filtre_por_otro_responsable(self):
        self._login("funcionario")
        resp = self._get(responsable=self.otro.user.pk)
        self.assertEqual(self._numbers(resp), [])
        # y su selector de responsable ni siquiera se muestra
        self.assertFalse(resp.context["show_responsible_filter"])

    def test_funcionario_solo_ve_las_propias_con_cualquier_orden_y_filtro(self):
        self._login("funcionario")
        resp = self._get(orden="-fecha", tipo=self.tipo_b.pk)
        self.assertEqual(self._numbers(resp), [self.a3.number])

    def test_opciones_de_responsable_solo_del_alcance(self):
        self._login("coordinador")      # activa DEL01: funcionario y otro
        labels = [o["label"] for o in self._get().context["responsible_options"]]
        self.assertEqual(sorted(labels), sorted(["Ana Alvarez", "Felipe Funciona"]))

    # --- tamaño de página y paginador ------------------------------------
    def test_tamano_de_pagina_por_defecto_es_5(self):
        self._login("superadmin")
        self._bulk(8)                                   # total 12
        resp = self._get()
        self.assertEqual(len(resp.context["rows"]), 5)
        self.assertContains(resp, "Mostrando 1 a 5 de 12 registros")

    def test_por_pagina_10_25_50(self):
        self._login("superadmin")
        self._bulk(30)                                  # total 34
        for size, expected in ((10, 10), (25, 25), (50, 34)):
            resp = self._get(por_pagina=size)
            self.assertEqual(len(resp.context["rows"]), expected, size)
        self.assertContains(self._get(por_pagina=10, page=2), "Mostrando 11 a 20 de 34 registros")

    def test_por_pagina_invalido_usa_5(self):
        self._login("superadmin")
        self._bulk(8)
        for bad in ("7", "0", "-5", "abc", "1000"):
            self.assertEqual(len(self._get(por_pagina=bad).context["rows"]), 5, bad)

    def test_enlaces_del_paginador_conservan_filtros_orden_y_tamano(self):
        self._login("superadmin")
        self._bulk(30)
        resp = self._get(tipo=self.tipo_a.pk, orden="-fecha", por_pagina=10,
                         desde="2000-01-01", q="maría")
        urls = [link["url"] for link in resp.context["page_links"] if not link.get("ellipsis")]
        urls += [resp.context["next_url"]]
        self.assertTrue(urls)
        for url in urls:
            for fragment in (f"tipo={self.tipo_a.pk}", "orden=-fecha", "por_pagina=10",
                             "desde=2000-01-01"):
                self.assertIn(fragment, url)

    def test_paginador_con_puntos_suspensivos(self):
        self._login("superadmin")
        self._bulk(96)                                  # total 100 -> 20 páginas de 5
        resp = self._get(page=10)
        links = resp.context["page_links"]
        self.assertTrue(any(link.get("ellipsis") for link in links))
        self.assertLess(len(links), 20)
        self.assertContains(resp, "…")
        numbers = [link["number"] for link in links if not link.get("ellipsis")]
        self.assertIn(1, numbers)
        self.assertIn(20, numbers)
        self.assertIn(10, numbers)

    def test_pagina_fuera_de_rango_e_invalida(self):
        self._login("superadmin")
        self._bulk(8)
        self.assertContains(self._get(page=999), "Mostrando 11 a 12 de 12 registros")
        self.assertContains(self._get(page="abc"), "Mostrando 1 a 5 de 12 registros")

    def test_sin_resultados_con_filtros_ofrece_limpiar(self):
        self._login("superadmin")
        resp = self._get(desde="2030-01-01")
        self.assertEqual(resp.context["empty_kind"], "results")
        self.assertContains(resp, "No encontramos resultados")

    # --- rendimiento ------------------------------------------------------
    def test_consultas_no_crecen_con_la_cantidad_de_filas(self):
        self._login("superadmin")
        with CaptureQueriesContext(connection) as few:
            self._get(por_pagina=5)
        self._bulk(40)
        with CaptureQueriesContext(connection) as many:
            self._get(por_pagina=50)
        # Con select_related el costo por fila es 0: la diferencia es mínima.
        self.assertLessEqual(len(many), len(few) + 2)

    # --- exportación ------------------------------------------------------
    def test_exportacion_aplica_filtros_y_orden_del_listado(self):
        self._login("superadmin")
        self.assertEqual(
            self._export_numbers(tipo=self.tipo_a.pk, orden="-numero"),
            [self.a2.number, self.a1.number],
        )
        self.assertEqual(
            self._export_numbers(desde="2026-03-08", hasta="2026-03-16", orden="fecha"),
            [self.a1.number, self.a4.number],
        )
        self.assertEqual(self._export_numbers(responsable=self.otro.user.pk, orden="-numero"),
                         [self.a4.number, self.a2.number])

    def test_exportacion_ignora_la_paginacion_y_exporta_todo(self):
        self._login("superadmin")
        self._bulk(26)                                  # total 30
        numbers = self._export_numbers(por_pagina=5, page=2)
        self.assertEqual(len(numbers), 30)

    def test_exportacion_respeta_alcance_por_rol(self):
        self._login("funcionario")
        self.assertEqual(self._export_numbers(responsable=self.otro.user.pk), [])
        self.assertEqual(sorted(self._export_numbers()), sorted([self.a1.number, self.a3.number]))

    def test_exportacion_con_parametros_invalidos_no_falla(self):
        self._login("superadmin")
        numbers = self._export_numbers(tipo="x", orden="zzz", desde="mal", por_pagina="9")
        self.assertEqual(len(numbers), 4)

    def test_enlace_de_exportacion_lleva_filtros_pero_no_pagina(self):
        self._login("superadmin")
        resp = self._get(tipo=self.tipo_a.pk, orden="-fecha", por_pagina=10, page=1)
        url = resp.context["export_url"]
        self.assertIn(f"tipo={self.tipo_a.pk}", url)
        self.assertIn("orden=", url)
        self.assertNotIn("page=", url)
        self.assertNotIn("por_pagina", url)

    def test_exportacion_exige_permiso_y_sesion(self):
        resp = self.client.get(reverse("attentions_export"))
        self.assertEqual(resp.status_code, 302)

    # --- controles en la plantilla ---------------------------------------
    def test_la_plantilla_muestra_los_controles_nuevos(self):
        self._login("superadmin")
        html = self._get().content.decode()
        for element_id in ("tipo", "subtipo", "delegacion", "responsable",
                           "desde", "hasta", "por_pagina"):
            self.assertRegex(html, rf'id="{element_id}"', element_id)
        self.assertIn("attention_list.js", html)
        self.assertIn("attention_list.css", html)

    def test_orden_elegido_se_conserva_al_filtrar(self):
        self._login("superadmin")
        html = self._get(orden="-fecha").content.decode()
        self.assertRegex(html, r'<input type="hidden" name="orden" value="-fecha">')
