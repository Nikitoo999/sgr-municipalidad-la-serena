"""Pruebas de la Parte 4: motor CRUD aplicado a Delegaciones.

Cubre: permisos por rol (solo el superadmin entra), búsqueda sin tildes ni
mayúsculas y con varias palabras, filtro de estado, paginación conservando
filtros, exportación que respeta filtros y columnas, código duplicado, color
inválido, eliminación bloqueada con relaciones y borrado lógico sin ellas,
AuditLog, 403 por URL directa de otro rol y 404 inexistente.
"""
from datetime import date

from django.contrib.auth.models import User
from django.test import TestCase, override_settings
from django.urls import reverse
from django.utils import timezone

from openpyxl import load_workbook

from operaciones.models import Attention, AttentionSubtype, AttentionType, Employee, Neighbor
from portal.models import AuditLog
from portal.tests.utils import TEST_STORAGES
from sgr_core.models import Delegation, InstitutionalGoal, Period, Role, RolePermission
from sgr_core.roles import ensure_base_roles


@override_settings(STORAGES=TEST_STORAGES)
class CrudDelegationsTestCase(TestCase):
    @classmethod
    def setUpTestData(cls):
        ensure_base_roles(Role, RolePermission)
        cls.roles = {role.code: role for role in Role.objects.all()}

        cls.d1 = Delegation.objects.create(
            name="Centro", code="DEL01", commune="La Serena",
            address="Calle Principal 1", color="#C62828",
        )
        cls.d2 = Delegation.objects.create(
            name="Avenida del Mar", code="DEL02", commune="La Serena",
            address="Av. Costanera 2", color="#1D6FA5",
        )
        cls.d3 = Delegation.objects.create(
            name="Las Compañías", code="DEL03", commune="La Serena",
            address="Calle Secundaria 3", color="#6A1B9A",
        )
        cls.d4 = Delegation.objects.create(
            name="Sector Cerrado", code="DEL04", status="Inactiva",
        )

        cls.superuser = User.objects.create_user(
            username="superadmin", password="Demo2026!", is_superuser=True, is_staff=True,
        )

        cls._rut_seq = 0

        def _employee(username, role_code, delegation):
            user = User.objects.create_user(username=username, password="Demo2026!")
            cls._rut_seq += 1
            return Employee.objects.create(
                user=user, rut=f"11.111.11{cls._rut_seq}-K",
                delegation=delegation, role=cls.roles[role_code],
            )

        cls.coordinador = _employee("coordinador", "ROL02", cls.d1)
        cls.coordinador.authorized_delegations.set([cls.d1, cls.d2])
        cls.delegado = _employee("delegado", "ROL03", cls.d1)
        cls.funcionario = _employee("funcionario", "ROL04", cls.d1)

    # ---------------------------------------------------------------
    # Permisos por rol
    # ---------------------------------------------------------------

    def test_listado_solo_superadmin(self):
        self.client.login(username="superadmin", password="Demo2026!")
        self.assertEqual(self.client.get(reverse("delegations")).status_code, 200)

        for username in ("coordinador", "delegado", "funcionario"):
            self.client.logout()
            self.client.login(username=username, password="Demo2026!")
            resp = self.client.get(reverse("delegations"))
            self.assertEqual(resp.status_code, 403)
            self.assertContains(resp, "Acceso restringido", status_code=403)

    def test_anonimo_redirige_a_login(self):
        resp = self.client.get(reverse("delegations"))
        self.assertEqual(resp.status_code, 302)
        self.assertIn(reverse("login"), resp["Location"])

    def test_403_url_directa_de_otro_rol(self):
        self.client.login(username="delegado", password="Demo2026!")
        self.assertEqual(
            self.client.get(reverse("delegations_detail", args=[self.d1.pk])).status_code, 403
        )
        self.assertEqual(
            self.client.get(reverse("delegations_edit", args=[self.d1.pk])).status_code, 403
        )

    def test_404_inexistente(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations_detail", args=[99999]))
        self.assertEqual(resp.status_code, 404)
        self.assertEqual(
            self.client.get(reverse("delegations_edit", args=[99999])).status_code, 404
        )

    def test_eliminar_solo_post_y_confirmado(self):
        self.client.login(username="superadmin", password="Demo2026!")
        # GET en eliminar no es una pantalla: 404.
        self.assertEqual(
            self.client.get(reverse("delegations_delete", args=[self.d4.pk])).status_code, 404
        )
        resp = self.client.post(reverse("delegations_delete", args=[self.d4.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d4.refresh_from_db()
        self.assertIsNotNone(self.d4.deleted_at)

    # ---------------------------------------------------------------
    # CRUD completo como superadmin
    # ---------------------------------------------------------------

    def test_crear_consultar_editar(self):
        self.client.login(username="superadmin", password="Demo2026!")

        resp = self.client.post(reverse("delegations_new"), {
            "code": "del08", "name": "Nueva Delegación", "commune": "La Serena",
            "address": "Calle Nueva 8", "color": "#920000", "status": "Activa",
        })
        self.assertEqual(resp.status_code, 302)
        created = Delegation.alive.get(code="DEL08")  # se guarda en mayúsculas
        self.assertEqual(created.name, "Nueva Delegación")

        detail = self.client.get(reverse("delegations_detail", args=[created.pk]))
        self.assertEqual(detail.status_code, 200)
        self.assertContains(detail, "Consulta de registro")
        self.assertContains(detail, "disabled")

        resp = self.client.post(reverse("delegations_edit", args=[created.pk]), {
            "code": "DEL08", "name": "Nueva Delegación Editada",
            "commune": "La Serena", "address": "Calle Nueva 8",
            "color": "#17663b", "status": "Inactiva",
        })
        self.assertEqual(resp.status_code, 302)
        created.refresh_from_db()
        self.assertEqual(created.name, "Nueva Delegación Editada")
        self.assertEqual(created.status, "Inactiva")

    def test_codigo_duplicado_rechazado(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.post(reverse("delegations_new"), {
            "code": "DEL01", "name": "Otra", "commune": "La Serena",
            "address": "", "color": "", "status": "Activa",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "Ya existe una delegación con este código", status_code=200)
        self.assertFalse(Delegation.alive.filter(name="Otra").exists())

    def test_color_invalido_rechazado(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.post(reverse("delegations_new"), {
            "code": "DEL09", "name": "Con color malo", "commune": "La Serena",
            "address": "", "color": "rojo", "status": "Activa",
        })
        self.assertEqual(resp.status_code, 200)
        self.assertContains(resp, "#RRGGBB", status_code=200)
        self.assertFalse(Delegation.alive.filter(code="DEL09").exists())

    # ---------------------------------------------------------------
    # Búsqueda, filtro y paginación
    # ---------------------------------------------------------------

    def _codes_shown(self, resp):
        """Códigos de la primera columna (Código) de las filas mostradas."""
        return [row["cells"][0]["text"] for row in resp.context["rows"]]

    def _flashed_text(self, resp):
        """Texto de los messages de Django (solo SweetAlert2 los muestra).

        Se lee ``sgr_messages`` (lista ya serializada por el context
        processor) porque el storage de ``messages`` queda consumido al
        serializarlo en el render.
        """
        return " ".join(item["text"] for item in resp.context["sgr_messages"])

    def test_busqueda_sin_tildes_ni_mayusculas_y_varias_palabras(self):
        self.client.login(username="superadmin", password="Demo2026!")
        for query in ("centro", "CENTRO", "Centro"):
            resp = self.client.get(reverse("delegations"), {"q": query})
            self.assertEqual(self._codes_shown(resp), ["DEL01"])
        # "Las Compañías" se encuentra sin tilde y en minúsculas.
        for query in ("companias", "COMPANIAS", "las companias"):
            resp = self.client.get(reverse("delegations"), {"q": query})
            self.assertEqual(self._codes_shown(resp), ["DEL03"])

    def test_filtro_estado(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations"), {"estado": "Inactiva"})
        self.assertEqual(self._codes_shown(resp), ["DEL04"])
        resp = self.client.get(reverse("delegations"), {"estado": "Activa"})
        self.assertEqual(self._codes_shown(resp), ["DEL01", "DEL02", "DEL03"])

    def test_paginacion_conserva_filtros(self):
        for i in range(5, 12):
            Delegation.objects.create(
                name=f"Extra {i}", code=f"DEL{i:02d}", commune="La Serena",
            )
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations"), {"q": "extra", "estado": "Activa"})
        self.assertContains(resp, "Mostrando 1 a 5 de 7 registros")
        # La paginación conserva la búsqueda y el filtro en la URL.
        self.assertContains(resp, "q=extra")
        self.assertContains(resp, "estado=Activa")
        # Página 2 sigue mostrando resultados del filtro.
        resp2 = self.client.get(reverse("delegations"), {"q": "extra", "estado": "Activa", "page": 2})
        self.assertContains(resp2, "Mostrando 6 a 7 de 7 registros")
        self.assertContains(resp2, "q=extra")

    def test_estado_vacio_distinto(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations"), {"q": "noexiste"})
        self.assertContains(resp, "No encontramos resultados")
        self.assertNotContains(resp, "Aún no hay registros")
        # Sin ningún registro vivo: estado "aún no hay registros".
        Delegation.objects.update(deleted_at=timezone.now())
        resp = self.client.get(reverse("delegations"))
        self.assertContains(resp, "Aún no hay registros")

    # ---------------------------------------------------------------
    # Exportación
    # ---------------------------------------------------------------

    def test_exportacion_respeta_filtros_y_columnas(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations_export"), {"estado": "Inactiva"})
        self.assertEqual(resp.status_code, 200)
        self.assertEqual(
            resp["Content-Type"],
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        self.assertIn("delegaciones_", resp["Content-Disposition"])

        import io

        workbook = load_workbook(io.BytesIO(resp.content))
        sheet = workbook.active
        headers = [cell.value for cell in sheet[1]]
        self.assertEqual(headers, ["Código", "Nombre", "Comuna", "Dirección", "Estado"])
        self.assertNotIn("Color", headers)  # nunca columnas no configuradas
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 1)  # solo la inactiva
        self.assertEqual(rows[0][0], "DEL04")

    def test_exportacion_busca_todas_las_filas_filtradas(self):
        for i in range(5, 12):
            Delegation.objects.create(name=f"Extra {i}", code=f"DEL{i:02d}")
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations_export"), {"q": "extra"})
        import io

        workbook = load_workbook(io.BytesIO(resp.content))
        sheet = workbook.active
        self.assertEqual(sheet.max_row - 1, 7)  # todas las filtradas, no solo la página

    def test_exportar_sin_permiso_403(self):
        self.client.login(username="coordinador", password="Demo2026!")
        resp = self.client.get(reverse("delegations_export"))
        self.assertEqual(resp.status_code, 403)

    # ---------------------------------------------------------------
    # Eliminación: regla de relaciones y borrado lógico
    # ---------------------------------------------------------------

    def _crear_atencion(self, delegation):
        atype = AttentionType.objects.create(code="TIP01", name="Reclamo")
        subtype = AttentionSubtype.objects.create(
            code="SUB01", name="Ruido", attention_type=atype,
        )
        neighbor = Neighbor.objects.create(
            identifier="11.111.111-1", full_name="Vecino Demo",
        )
        return Attention.objects.create(
            neighbor=neighbor, delegation=delegation, attention_type=atype,
            attention_subtype=subtype, description="Ruido en la vía pública",
            responsible=self.superuser,
        )

    def test_eliminar_bloqueado_con_relaciones(self):
        self.coordinador.authorized_delegations.clear()
        self.client.login(username="superadmin", password="Demo2026!")
        # DEL01 tiene empleados (coordinador, delegado, funcionario).
        resp = self.client.post(reverse("delegations_delete", args=[self.d1.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d1.refresh_from_db()
        self.assertIsNone(self.d1.deleted_at)
        follow = self.client.get(reverse("delegations"))
        self.assertIn("Márcala como Inactiva", self._flashed_text(follow))

        # DEL02 queda libre: se borra de forma lógica.
        resp = self.client.post(reverse("delegations_delete", args=[self.d2.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d2.refresh_from_db()
        self.assertIsNotNone(self.d2.deleted_at)
        self.assertNotIn("DEL02", self._codes_shown(self.client.get(reverse("delegations"))))

    def test_eliminar_bloqueado_con_delegacion_autorizada(self):
        self.client.login(username="superadmin", password="Demo2026!")
        # El coordinador tiene DEL02 en sus delegaciones autorizadas (M2M).
        resp = self.client.post(reverse("delegations_delete", args=[self.d2.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d2.refresh_from_db()
        self.assertIsNone(self.d2.deleted_at)
        self.assertIn("usuario(s)", self._flashed_text(self.client.get(reverse("delegations"))))

    def test_eliminar_bloqueado_con_atencion(self):
        self._crear_atencion(self.d3)
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.post(reverse("delegations_delete", args=[self.d3.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d3.refresh_from_db()
        self.assertIsNone(self.d3.deleted_at)
        # El mensaje viaja al listado (solo SweetAlert2 lo muestra).
        follow = self.client.get(reverse("delegations"))
        self.assertIn("No se puede eliminar", self._flashed_text(follow))

    def test_eliminar_bloqueado_con_meta(self):
        period = Period.objects.create(
            name="2026-T1", quarter=1,
            start_date=date(2026, 1, 1), end_date=date(2026, 3, 31),
        )
        InstitutionalGoal.objects.create(
            name="Meta demo", period=period, delegation=self.d3,
        )
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.post(reverse("delegations_delete", args=[self.d3.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d3.refresh_from_db()
        self.assertIsNone(self.d3.deleted_at)
        self.assertIn("meta(s)", self._flashed_text(self.client.get(reverse("delegations"))))

    def test_borrado_logico_sin_relaciones(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.post(reverse("delegations_delete", args=[self.d4.pk]))
        self.assertEqual(resp.status_code, 302)
        self.d4.refresh_from_db()
        self.assertIsNotNone(self.d4.deleted_at)
        self.assertFalse(Delegation.alive.filter(pk=self.d4.pk).exists())
        # El registro ya no es accesible (404).
        self.assertEqual(
            self.client.get(reverse("delegations_detail", args=[self.d4.pk])).status_code, 404
        )

    # ---------------------------------------------------------------
    # Auditoría
    # ---------------------------------------------------------------

    def test_auditlog_registra_acciones(self):
        self.client.login(username="superadmin", password="Demo2026!")
        self.client.post(reverse("delegations_new"), {
            "code": "DEL20", "name": "Auditoría Demo", "commune": "La Serena",
            "address": "", "color": "", "status": "Activa",
        })
        created = Delegation.alive.get(code="DEL20")
        self.client.post(reverse("delegations_edit", args=[created.pk]), {
            "code": "DEL20", "name": "Auditoría Demo 2", "commune": "La Serena",
            "address": "", "color": "", "status": "Activa",
        })
        self.client.get(reverse("delegations_export"))
        self.client.post(reverse("delegations_delete", args=[created.pk]))

        actions = set(
            AuditLog.objects.filter(module="delegations").values_list("action", flat=True)
        )
        self.assertIn("create", actions)
        self.assertIn("edit", actions)
        self.assertIn("export", actions)
        self.assertIn("delete", actions)

    def test_auditlog_403_otro_rol(self):
        self.client.login(username="delegado", password="Demo2026!")
        self.client.get(reverse("delegations"))
        self.assertTrue(
            AuditLog.objects.filter(action="access_denied", module="delegations").exists()
        )




    def test_toolbar_sin_boton_buscar(self):
        self.client.login(username="superadmin", password="Demo2026!")
        resp = self.client.get(reverse("delegations"))
        self.assertEqual(resp.status_code, 200)
        # Búsqueda inteligente: el campo q existe y no hay botón Buscar.
        self.assertContains(resp, 'id="q"')
        self.assertNotContains(resp, 'id="searchBtn"')

