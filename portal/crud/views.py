"""Vistas base del motor CRUD (Parte 4).

Todas las vistas de mantenedor (lista, crear, ver, editar, eliminar, exportar)
viven aquí y se configuran solo con ``module_key`` + la acción de permiso.
Reglas comunes:

- Toda autorización se valida en el servidor (``ModulePermissionMixin`` +
  alcance del queryset): ocultar botones no cuenta.
- El queryset excluye registros con borrado lógico (``deleted_at`` no nulo) vía
  el manager ``alive`` del modelo.
- Objeto existente pero fuera de alcance: 403 + AuditLog. Inexistente: 404.
- Las acciones que modifican datos van solo por POST (CSRF incluido).
"""

from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.core.exceptions import PermissionDenied
from django.core.paginator import Paginator
from django.db.models import Count
from django.http import Http404, HttpResponse
from django.shortcuts import redirect
from django.urls import reverse
from django.utils import timezone
from django.views import View
from django.views.generic import TemplateView
from django.views.generic.edit import FormView

from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from ..access import ModulePermissionMixin, get_access
from ..audit import audit
from ..context_processors import badge_colors
from .config import CELL_DELEGATION, CELL_STATUS, ModuleConfig, get_module
from .search import apply_search

#: Encabezado del Excel con el color institucional.
HEADER_FILL = PatternFill("solid", fgColor="920000")
HEADER_FONT = Font(color="FFFFFF", bold=True)
HEADER_ALIGN = Alignment(horizontal="center", vertical="center")
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)


def _display(value):
    """Texto de una celda para listado/Excel (nunca datos sensibles)."""
    if value is None or value == "":
        return "—"
    return str(value)


def _export_value(obj, column):
    """Valor plano de una columna para el Excel (nunca datos sensibles)."""
    value = getattr(obj, column.field_name, None)
    if hasattr(value, "all"):  # relación (manager/queryset)
        value = list(value.all())
    if isinstance(value, (list, tuple)):
        parts = []
        for item in value:
            code = getattr(item, "code", "")
            name = getattr(item, "name", "")
            if code and name:
                parts.append(f"{code} · {name}")
            else:
                parts.append(str(item))
        value = ", ".join(parts)
    elif callable(value):
        value = value()
    return _display(value)


def _delegation_chips(value):
    """Normaliza una delegación, un conjunto o una lista a objetos."""
    if value is None:
        return []
    if hasattr(value, "all"):  # manager / queryset (varios chips)
        return list(value.all())
    if isinstance(value, (list, tuple)):
        return list(value)
    return [value]


def resolve_cell(obj, column):
    """Valor listo para la plantilla de una celda del listado."""
    value = getattr(obj, column.field_name, None)
    if column.cell == CELL_STATUS:
        return {"kind": "status", "text": str(value or "")}
    if column.cell == CELL_DELEGATION:
        chips = []
        for delegation in _delegation_chips(value):
            dc, dm = badge_colors(delegation.color)
            chips.append(
                {
                    "code": delegation.code,
                    "name": delegation.name,
                    "color": delegation.color or "",
                    "dc": dc,
                    "dm": dm,
                }
            )
        return {"kind": "delegation", "chips": chips}
    return {"kind": "text", "text": _display(value)}


def resolve_object_or_deny(request, config, queryset, pk):
    """Objeto del queryset acotado, con 403/404 según el caso.

    - Existe en el modelo vivo pero fuera del alcance -> 403 (+ auditoría).
    - No existe -> 404.
    """
    try:
        obj = queryset.filter(pk=pk).first()
    except (ValueError, TypeError):
        raise Http404("Registro inexistente.")
    if obj is not None:
        return obj
    if config.model.alive.filter(pk=pk).exists():
        audit(
            request, "access_denied", module=config.module_key,
            object_repr=str(pk),
            detail=f"{request.method} {request.path} (pk={pk} fuera de alcance)",
        )
        raise PermissionDenied
    raise Http404("Registro inexistente.")


def scoped_queryset(request, config):
    """Queryset vivo del módulo, con el orden y el alcance aplicados."""
    queryset = config.model.alive.all()
    if config.ordering:
        queryset = queryset.order_by(*config.ordering)
    if config.scope is not None:
        queryset = config.scope(request, queryset)
    return queryset


def module_filters(request, config):
    """Selects de la grilla de filtros del módulo y sus valores vigentes.

    Devuelve ``(specs, values)``:

    - ``specs``: uno por filtro declarado (``FilterSpec``), con las opciones
      que puede elegir el usuario (el callable del módulo ya entrega las del
      alcance de su rol), el valor vigente y su etiqueta para la píldora.
    - ``values``: ``{clave: valor}`` con los valores aceptados. Un valor GET
      que no esté entre las opciones del filtro se ignora (no se filtra por
      él), igual que ocurre con un ``estado`` desconocido.
    """
    specs = []
    values = {}
    for spec in config.filters:
        options = [
            (str(value), str(label))
            for value, label in (spec.options(request) if spec.options else ())
        ]
        labels = dict(options)
        raw = (request.GET.get(spec.key) or "").strip()
        value = raw if raw in labels else ""
        values[spec.key] = value
        specs.append(
            {
                "key": spec.key,
                "label": spec.label,
                "empty_label": spec.empty_label,
                "value": value,
                "text": labels.get(value, value),
                "options": [
                    {
                        "value": option_value,
                        "label": label,
                        "selected": option_value == value,
                    }
                    for option_value, label in options
                ],
            }
        )
    return specs, values


def _filter_matches(obj, spec, value):
    """Compara un registro con el valor del filtro (modo lista de SQLite).

    En claves foráneas compara el ``pk`` del relacionado (que es el valor de
    la opción); en campos planos compara el valor como texto.
    """
    resolved = obj
    for part in spec.field_name.split("__"):
        if resolved is None:
            return False
        resolved = getattr(resolved, part, None)
    if resolved is None:
        return False
    pk = getattr(resolved, "pk", None)
    return str(pk if pk is not None else resolved) == str(value)


def apply_module_filters(rows, config, values):
    """Aplica los filtros declarados por el módulo al resultado del listado.

    ``rows`` puede ser un queryset (MariaDB, o SQLite sin búsqueda) o la lista
    que devuelve ``apply_search`` en SQLite: en ambos casos el resultado es el
    mismo y se conserva el orden del listado.
    """
    active = [
        (spec, values[spec.key]) for spec in config.filters
        if values.get(spec.key)
    ]
    if not active:
        return rows
    if hasattr(rows, "filter"):
        for spec, value in active:
            rows = rows.filter(**{spec.lookup_path: value})
        return rows
    return [
        obj for obj in rows
        if all(_filter_matches(obj, spec, value) for spec, value in active)
    ]


def filtered_queryset(request, config):
    """Queryset vivo + alcance + filtros del módulo y de estado (sin búsqueda)."""
    queryset = scoped_queryset(request, config)
    _, values = module_filters(request, config)
    queryset = apply_module_filters(queryset, config, values)
    estado = request.GET.get("estado", "")
    valid = [option[0] for option in config.status_options]
    if config.status_field and estado in valid:
        queryset = queryset.filter(**{config.status_field: estado})
    return queryset


#: Filas por página que acepta el selector del pie del listado. Cualquier otro
#: valor pedido por GET se ignora y se usa el ``per_page`` del módulo.
PER_PAGE_CHOICES = (5, 10, 20)


def per_page_value(request, config):
    """Filas por página efectivas del listado (solo 5, 10 o 20)."""
    try:
        value = int((request.GET.get("per_page") or "").strip())
    except ValueError:
        return config.per_page
    return value if value in PER_PAGE_CHOICES else config.per_page


def apply_status_filter(rows, config, estado):
    """Aplica el filtro de estado al resultado de la búsqueda.

    ``rows`` puede ser un queryset (MariaDB, o SQLite sin búsqueda) o la lista
    que devuelve ``apply_search`` en SQLite: en ambos casos se conserva el
    orden y el resultado es el mismo.
    """
    valid = [option[0] for option in config.status_options]
    if not config.status_field or estado not in valid:
        return rows
    if hasattr(rows, "filter"):
        return rows.filter(**{config.status_field: estado})
    return [
        obj for obj in rows
        if getattr(obj, config.status_field, None) == estado
    ]


def status_counts(config, rows):
    """Contadores por estado para las píldoras del listado.

    Recibe el resultado con alcance y búsqueda (sin el filtro de estado) y
    devuelve el total y el conteo de cada opción de ``status_options``. Con
    queryset se resuelve en una sola consulta agregada; con la lista ya
    materializada de SQLite se cuenta en memoria, sin consultas extra.
    """
    field, options = config.status_field, config.status_options
    if not field or not options:
        return {"all": 0, "options": []}
    if hasattr(rows, "values"):
        grouped = rows.order_by().values(field).annotate(total=Count("pk"))
        by_value = {row[field]: row["total"] for row in grouped}
    else:
        by_value = {}
        for obj in rows:
            value = getattr(obj, field, None)
            by_value[value] = by_value.get(value, 0) + 1
    return {
        "all": sum(by_value.values()),
        "options": [
            {"value": value, "label": label, "count": by_value.get(value, 0)}
            for value, label in options
        ],
    }


def querystring(query, estado, extra=None):
    """URL query conservando búsqueda y filtros (para la paginación)."""
    params = {}
    if query:
        params["q"] = query
    if estado:
        params["estado"] = estado
    if extra:
        params.update({key: value for key, value in extra.items() if value})
    return f"?{urlencode(params)}" if params else ""


#: Textos del modo del formulario (migas y título).
_PAGE_ACTIONS = {
    "create": "Crear registro",
    "edit": "Editar registro",
    "view": "Consulta de registro",
}


class CrudBaseMixin:
    """Resuelve la configuración del módulo y expone los permisos.

    ``setup`` corre antes de los ``dispatch`` de login y de permisos, así
    ``self.module`` ya existe cuando ``ModulePermissionMixin`` valida el
    permiso del módulo/acción.
    """

    module_key = None
    action = "view"
    config = None

    def setup(self, request, *args, **kwargs):
        super().setup(request, *args, **kwargs)
        self.config = get_module(self.module_key)
        self.module = self.config.module_key

    def access_flags(self):
        """Permisos de acción del usuario sobre este módulo (para el template)."""
        info = get_access(self.request)
        return {
            "can_view": info.can(self.module, "view"),
            "can_create": info.can(self.module, "create"),
            "can_edit": info.can(self.module, "edit"),
            "can_delete": info.can(self.module, "delete"),
            "can_export": info.can(self.module, "export"),
        }



class CrudListView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, TemplateView):
    """Listado con búsqueda, píldoras de estado, filtros del módulo y paginación.

    Además de ``q`` y ``estado`` acepta ``per_page`` (5, 10 o 20; cualquier
    otro valor usa el ``per_page`` del módulo) y las claves de los filtros que
    el módulo declare en ``config.filters`` (nunca las reservadas). Expone
    ``status_counts`` con los contadores por estado para las píldoras y
    ``filters`` con los selects de la grilla de filtros.
    """

    template_name = "portal/crud/list.html"
    action = "view"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        config = self.config
        request = self.request
        query = (request.GET.get("q") or "").strip()
        estado = request.GET.get("estado") or ""
        per_page = per_page_value(request, config)
        filter_specs, filter_values = module_filters(request, config)
        filter_params = {
            key: value for key, value in filter_values.items() if value
        }

        # Alcance + filtros del módulo + búsqueda (sin estado): esta base
        # alimenta las píldoras de estado y el filtro de estado se aplica
        # encima, así el buscador y los contadores siempre se refieren al
        # mismo conjunto.
        base_rows = apply_module_filters(
            apply_search(
                scoped_queryset(request, config), config.search_fields, query,
            ),
            config, filter_values,
        )
        rows_all = apply_status_filter(base_rows, config, estado)
        paginator = Paginator(rows_all, per_page)
        page_obj = paginator.get_page(request.GET.get("page"))

        rows = [
            {
                "pk": obj.pk,
                "cells": [resolve_cell(obj, column) for column in config.columns],
            }
            for obj in page_obj.object_list
        ]

        # Píldoras de estado: conservan q, filtros del módulo y per_page.
        counts = status_counts(config, base_rows)
        counts["all_url"] = querystring(
            query, "", {**filter_params, "per_page": per_page},
        )
        for pill in counts["options"]:
            pill["url"] = querystring(
                query, pill["value"], {**filter_params, "per_page": per_page},
            )
            pill["active"] = pill["value"] == estado

        # Etiquetas de filtros activos: cada una quita solo su propio filtro.
        status_labels = dict(config.status_options)
        active_filters = []
        if query:
            active_filters.append({
                "kind": "q",
                "text": f"Búsqueda: «{query}»",
                "url": querystring(
                    "", estado, {**filter_params, "per_page": per_page},
                ),
            })
        if estado:
            active_filters.append({
                "kind": "estado",
                "text": f"Estado: {status_labels.get(estado, estado)}",
                "url": querystring(
                    query, "", {**filter_params, "per_page": per_page},
                ),
            })
        for spec in filter_specs:
            if not spec["value"]:
                continue
            others = {
                key: value for key, value in filter_params.items()
                if key != spec["key"]
            }
            active_filters.append({
                "kind": spec["key"],
                "text": f"{spec['label']}: {spec['text']}",
                "url": querystring(
                    query, estado, {**others, "per_page": per_page},
                ),
            })

        page_links = [
            {
                "number": number,
                "url": querystring(
                    query, estado,
                    {**filter_params, "per_page": per_page, "page": number},
                ),
                "current": number == page_obj.number,
            }
            for number in range(1, paginator.num_pages + 1)
        ]

        context.update(
            {
                "active_module": config.module_key,
                "config": config,
                "rows": rows,
                "columns": config.columns,
                "page_obj": page_obj,
                "page_links": page_links,
                "prev_url": querystring(query, estado, {**filter_params, "per_page": per_page, "page": page_obj.previous_page_number() if page_obj.has_previous() else None}),
                "next_url": querystring(query, estado, {**filter_params, "per_page": per_page, "page": page_obj.next_page_number() if page_obj.has_next() else None}),
                "showing_from": page_obj.start_index(),
                "showing_to": page_obj.end_index(),
                "total": paginator.count,
                "query": query,
                "estado": estado,
                "status_options": config.status_options,
                "status_counts": counts,
                "active_filters": active_filters,
                "filters": filter_specs,
                "filter_params": [
                    {"name": key, "value": value}
                    for key, value in filter_params.items()
                ],
                "per_page": per_page,
                "per_page_options": PER_PAGE_CHOICES,
                "empty_kind": "results" if (query or estado or filter_params) else "records",
                "search_url": reverse(config.module_key),
                "export_url": f"{reverse(f'{config.module_key}_export')}{querystring(query, estado, filter_params)}",
                "new_url": reverse(f"{config.module_key}_new"),
                "detail_url_name": f"{config.module_key}_detail",
                "edit_url_name": f"{config.module_key}_edit",
                "delete_url_name": f"{config.module_key}_delete",
                **self.access_flags(),
            }
        )
        return context



class CrudFormView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, FormView):
    """Base de crear/editar/consultar: formulario compartido."""

    template_name = "portal/crud/form.html"
    action = "create"
    mode = "create"
    object = None

    def get_template_names(self):
        """Plantilla del módulo (por defecto ``portal/crud/form.html``)."""
        return [self.config.form_template]

    def get_form_class(self):
        return self.config.form_class

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        if self.object is not None:
            kwargs["instance"] = self.object
        # El request permite al formulario resolver reglas que dependen del
        # operador (roles ofrecidos, delegaciones administrables).
        kwargs["request"] = self.request
        return kwargs

    def get_success_url(self):
        return reverse(self.config.module_key)

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(
            {
                "active_module": self.config.module_key,
                "config": self.config,
                "mode": self.mode,
                "page_action": _PAGE_ACTIONS[self.mode],
                "has_errors": bool(getattr(context.get("form"), "errors", False)),
                "cancel_url": reverse(self.config.module_key),
                "is_create": self.mode == "create",
                "is_view": self.mode == "view",
                **self.access_flags(),
            }
        )
        return context


class CrudCreateView(CrudFormView):
    """Crear un registro (acción create, solo POST guarda)."""

    action = "create"
    mode = "create"

    def form_valid(self, form):
        obj = form.save()
        audit(
            self.request, "create", module=self.config.module_key,
            object_repr=str(obj), detail=f"Creado por {self.request.user}",
        )
        messages.success(
            self.request, f"{self.config.singular_label} creada correctamente."
        )
        return super().form_valid(form)


class CrudDetailView(CrudFormView):
    """Consultar un registro: formulario deshabilitado y botón Volver."""

    action = "view"
    mode = "view"

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.config.form_class(instance=self.object)
        for field in form.fields.values():
            field.widget.attrs["disabled"] = True
        return self.render_to_response(self.get_context_data(form=form))

    def post(self, request, *args, **kwargs):
        # Consultar no modifica datos: solo GET (evita guardados por error).
        raise Http404("Registro inexistente.")

    def get_object(self):
        queryset = scoped_queryset(self.request, self.config)
        return resolve_object_or_deny(
            self.request, self.config, queryset, self.kwargs["pk"]
        )



class CrudUpdateView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, FormView):
    """Editar un registro (acción edit)."""

    template_name = "portal/crud/form.html"
    action = "edit"
    mode = "edit"
    object = None

    def get_form_class(self):
        return self.config.form_class

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        return self.render_to_response(self.get_context_data(form=self.get_form()))

    def post(self, request, *args, **kwargs):
        self.object = self.get_object()
        return super().post(request, *args, **kwargs)

    def get_object(self):
        queryset = scoped_queryset(self.request, self.config)
        return resolve_object_or_deny(
            self.request, self.config, queryset, self.kwargs["pk"]
        )

    def get_form_kwargs(self):
        kwargs = super().get_form_kwargs()
        kwargs["instance"] = self.object
        # El formulario de Usuarios exige el operador (request) para resolver
        # roles y delegaciones administrables en el servidor.
        kwargs["request"] = self.request
        return kwargs

    def get_success_url(self):
        return reverse(self.config.module_key)

    def form_valid(self, form):
        obj = form.save()
        audit(
            self.request, "edit", module=self.config.module_key,
            object_repr=str(obj), detail=f"Editado por {self.request.user}",
        )
        messages.success(
            self.request, f"{self.config.singular_label} actualizada correctamente."
        )
        return redirect(self.get_success_url())

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(
            {
                "active_module": self.config.module_key,
                "config": self.config,
                "mode": self.mode,
                "page_action": _PAGE_ACTIONS[self.mode],
                "has_errors": bool(getattr(context.get("form"), "errors", False)),
                "cancel_url": reverse(self.config.module_key),
                "is_create": False,
                "is_view": False,
                **self.access_flags(),
            }
        )
        return context


class CrudDeleteView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, View):
    """Eliminar solo por POST: regla del módulo o borrado lógico."""

    action = "delete"

    def get(self, request, *args, **kwargs):
        # Eliminar no es una pantalla: la confirmación vive en el listado.
        raise Http404("Registro inexistente.")

    def post(self, request, *args, **kwargs):
        queryset = scoped_queryset(request, self.config)
        obj = resolve_object_or_deny(
            request, self.config, queryset, self.kwargs["pk"]
        )

        if self.config.delete_guard is not None:
            reason = self.config.delete_guard(obj, request)
            if reason:
                audit(
                    request, "delete", module=self.config.module_key,
                    object_repr=str(obj), detail=f"Bloqueado: {reason}",
                )
                messages.error(request, reason)
                return redirect(self.config.module_key)

        obj.deleted_at = timezone.now()
        obj.save(update_fields=["deleted_at", "updated_at"])
        audit(
            request, "delete", module=self.config.module_key,
            object_repr=str(obj), detail=f"Eliminado por {request.user}",
        )
        messages.success(
            request, f"{self.config.singular_label} eliminada correctamente."
        )
        return redirect(self.config.module_key)



class CrudExportView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, View):
    """Exporta a Excel TODAS las filas del resultado filtrado y con alcance."""

    action = "export"

    def get(self, request, *args, **kwargs):
        config = self.config
        query = (request.GET.get("q") or "").strip()
        estado = request.GET.get("estado") or ""

        queryset = filtered_queryset(request, config)
        rows_all = list(apply_search(queryset, config.search_fields, query))

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = config.list_title[:31]
        columns = config.export_columns

        for index, column in enumerate(columns, start=1):
            cell = sheet.cell(row=1, column=index, value=column.header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        for row_index, obj in enumerate(rows_all, start=2):
            for col_index, column in enumerate(columns, start=1):
                cell = sheet.cell(
                    row=row_index, column=col_index, value=_export_value(obj, column)
                )
                cell.border = BORDER

        for index, column in enumerate(columns, start=1):
            letter = get_column_letter(index)
            width = max(len(column.header) + 2, 14)
            for obj in rows_all:
                width = max(width, len(_export_value(obj, column)))
            sheet.column_dimensions[letter].width = min(width, 48)
        sheet.freeze_panes = "A2"

        filename = f"{config.slug}_{timezone.now():%Y%m%d}.xlsx"
        audit(
            request, "export", module=config.module_key,
            object_repr=f"{len(rows_all)} registros",
            detail=f"q={query!r} estado={estado!r}",
        )
        response = HttpResponse(
            content_type=(
                "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            ),
        )
        response["Content-Disposition"] = f'attachment; filename="{filename}"'
        workbook.save(response)
        return response




