"""Configuración del módulo de Atenciones (Parte 7).

El listado, la exportación y los formularios (crear/consultar/editar) se acotan
con :func:`attentions_scope`: superadmin según su delegación activa ("Todas" o
una), coordinador según su delegación activa, delegado su delegación fija y
funcionario únicamente donde él es responsable. Las vistas de crear/editar/ver
son propias para respetar el diseño "Ingreso de atención" (búsqueda de
solicitante, delegación fija por perfil, tipo/sub atención dependientes,
descripción con contador, estado editable solo al editar).
"""

from datetime import date, datetime, time, timedelta
from urllib.parse import urlencode

from django.contrib import messages
from django.contrib.auth.mixins import LoginRequiredMixin
from django.contrib.auth.models import User
from django.core.paginator import Paginator
from django.db.models import Case, IntegerField, Value, When
from django.http import HttpResponse, JsonResponse
from django.shortcuts import redirect
from django.urls import reverse
from django.utils import timezone
from django.views import View

from openpyxl.utils import get_column_letter
from openpyxl.workbook import Workbook

from operaciones.models import Attention, Neighbor
from sgr_core.models import AttentionSubtype, AttentionType

from ..access import (
    ALL_DELEGATIONS,
    ModulePermissionMixin,
    allowed_delegations,
    get_active_delegation,
    get_role,
)
from ..audit import audit
from .config import CELL_DELEGATION, CELL_STATUS, ColumnSpec, ModuleConfig, register_module
from .forms import AttentionForm
from .search import apply_search
from .views import (
    BORDER,
    HEADER_ALIGN,
    HEADER_FILL,
    HEADER_FONT,
    CrudBaseMixin,
    CrudCreateView,
    CrudDetailView,
    CrudListView,
    CrudUpdateView,
    filtered_queryset,
    resolve_cell,
    scoped_queryset,
)


def attentions_scope(request, queryset):
    """Atenciones visibles según rol + delegación activa.

    - Superadmin: todas, o solo las de su delegación activa.
    - Coordinador: su delegación activa (entre las autorizadas).
    - Delegado: su delegación fija.
    - Funcionario: solo las que él ingresó (responsable = usuario).
    """
    role = get_role(request.user)
    if role is None:
        return queryset.none()
    active = get_active_delegation(request)
    if active is not None and active != ALL_DELEGATIONS:
        queryset = queryset.filter(delegation_id=active.pk)
    if role.scope == "own_records":
        queryset = queryset.filter(responsible_id=request.user.pk)
    return queryset


def _deleg_label(delegation):
    return f"{delegation.code} | Delegación {delegation.name}"


def scope_note(request):
    """Aviso de alcance sobre la tabla (textos del diseño ``bannerFor``)."""
    role = get_role(request.user)
    if role is None:
        return ""
    active = get_active_delegation(request)
    if role.scope == "own_records":
        return "Tu perfil muestra solo las atenciones que ingresaste tú."
    if role.scope == "all":
        if active == ALL_DELEGATIONS:
            return "Vista global: todas las delegaciones."
        return (
            f"Filtrado por {_deleg_label(active)}. Cambia la delegación activa "
            "en la cabecera."
        )
    if role.scope == "authorized":
        count = len(allowed_delegations(request.user))
        return (
            f"Delegación activa: {_deleg_label(active)}. Puedes cambiarla en la "
            f"cabecera entre tus {count} delegaciones autorizadas."
        )
    return f"Tu perfil muestra solo registros de {_deleg_label(active)}."


ATTENTIONS_MODULE = ModuleConfig(
    module_key="attentions",
    slug="atenciones",
    model=Attention,
    form_class=AttentionForm,
    list_title="Atenciones",
    singular_label="Atención",
    columns=(
        ColumnSpec("number", "N°"),
        ColumnSpec("created_at_display", "Fecha y hora"),
        ColumnSpec("solicitante_name", "Solicitante"),
        ColumnSpec("type_name", "Tipo"),
        ColumnSpec("subtype_name", "Sub atención"),
        ColumnSpec("delegation", "Delegación", cell=CELL_DELEGATION),
        ColumnSpec("responsible_name", "Responsable"),
        ColumnSpec("status_display", "Estado", cell=CELL_STATUS),
    ),
    search_fields=(
        "number", "neighbor__full_name", "neighbor__identifier",
        "attention_type__name", "attention_subtype__name",
        "responsible__first_name", "responsible__last_name",
        "responsible__username", "delegation__name", "delegation__code",
    ),
    status_field="status",
    status_options=(
        ("received", "Ingresada"),
        ("in_progress", "En gestión"),
        ("resolved", "Resuelta"),
    ),
    ordering=("number",),
    scope=attentions_scope,
    form_template="portal/crud/attention_form.html",
)

register_module(ATTENTIONS_MODULE)


# ---------------------------------------------------------------------------
# Listado avanzado (filtros, orden y tamaño de página)
# ---------------------------------------------------------------------------
#
# Los mismos parámetros GET alimentan el listado y la exportación a Excel, así
# el archivo descargado siempre coincide con lo que se ve en pantalla (menos la
# paginación). Todo valor que llega por la URL se valida contra una lista
# blanca o se convierte a número/fecha; un valor inválido se ignora y nunca
# produce un error 500. El alcance por rol (``attentions_scope``) se aplica
# antes que cualquier filtro, así ningún parámetro puede ampliar lo visible.

#: Tamaños de página permitidos (el primero es el valor por defecto).
PAGE_SIZE_OPTIONS = (5, 10, 25, 50)
DEFAULT_PAGE_SIZE = PAGE_SIZE_OPTIONS[0]

#: Clave pública de orden -> campos reales de la base de datos. La clave
#: ``estado`` se resuelve aparte para ordenar por avance (no alfabético).
SORT_FIELDS = {
    "numero": ("number",),
    "fecha": ("created_at",),
    "solicitante": ("neighbor__full_name",),
    "tipo": ("attention_type__name",),
    "subtipo": ("attention_subtype__name",),
    "delegacion": ("delegation__code",),
    "responsable": (
        "responsible__first_name", "responsible__last_name",
        "responsible__username",
    ),
    "estado": ("_estado_orden",),
}
DEFAULT_SORT = "numero"

#: Campo de cada columna del listado -> clave de orden.
COLUMN_SORT_KEY = {
    "number": "numero",
    "created_at_display": "fecha",
    "solicitante_name": "solicitante",
    "type_name": "tipo",
    "subtype_name": "subtipo",
    "delegation": "delegacion",
    "responsible_name": "responsable",
    "status_display": "estado",
}

#: Orden lógico del avance de una atención (Ingresada -> En gestión -> Resuelta).
_STATUS_PROGRESS = (
    Attention.AttentionStatus.RECEIVED,
    Attention.AttentionStatus.IN_PROGRESS,
    Attention.AttentionStatus.RESOLVED,
)

#: Parámetros GET que se conservan al paginar, ordenar y exportar.
LIST_PARAM_KEYS = (
    "q", "estado", "tipo", "subtipo", "delegacion", "responsable",
    "desde", "hasta", "orden", "por_pagina",
)


def split_order(orden):
    """``"-fecha"`` -> ``("fecha", True)``; solo se admite UN guion inicial."""
    orden = orden or ""
    if orden.startswith("-"):
        return orden[1:], True
    return orden, False


def _parse_pk(raw):
    """Entero positivo o ``None`` (cualquier otra cosa se ignora)."""
    try:
        value = int(str(raw).strip())
    except (TypeError, ValueError):
        return None
    return value if value > 0 else None


def _parse_date(raw):
    """Fecha ISO (AAAA-MM-DD) o ``None``."""
    try:
        return date.fromisoformat(str(raw).strip())
    except (TypeError, ValueError):
        return None


def parse_list_params(request):
    """Lee y valida los parámetros del listado/exportación de atenciones.

    Devuelve un diccionario con los valores ya normalizados (cadenas vacías
    cuando el parámetro no viene o es inválido) más ``date_error`` si el rango
    de fechas es incoherente.
    """
    get = request.GET
    valid_states = {value for value, _label in ATTENTIONS_MODULE.status_options}

    estado = get.get("estado") or ""
    tipo = _parse_pk(get.get("tipo"))
    subtipo = _parse_pk(get.get("subtipo"))
    delegacion = _parse_pk(get.get("delegacion"))
    responsable = _parse_pk(get.get("responsable"))

    desde = _parse_date(get.get("desde"))
    hasta = _parse_date(get.get("hasta"))
    date_error = ""
    if desde and hasta and desde > hasta:
        date_error = "La fecha «Desde» no puede ser posterior a «Hasta»: se ignoró el rango."
        desde = hasta = None

    orden = (get.get("orden") or "").strip()
    if split_order(orden)[0] not in SORT_FIELDS:
        orden = ""

    por_pagina = _parse_pk(get.get("por_pagina"))
    if por_pagina not in PAGE_SIZE_OPTIONS:
        por_pagina = None

    return {
        "q": (get.get("q") or "").strip(),
        "estado": estado if estado in valid_states else "",
        "tipo": str(tipo) if tipo else "",
        "subtipo": str(subtipo) if subtipo else "",
        "delegacion": str(delegacion) if delegacion else "",
        "responsable": str(responsable) if responsable else "",
        "desde": desde.isoformat() if desde else "",
        "hasta": hasta.isoformat() if hasta else "",
        "orden": orden,
        "por_pagina": str(por_pagina) if por_pagina else "",
        "date_error": date_error,
    }


def _day_start(day):
    return timezone.make_aware(datetime.combine(day, time.min))


def _apply_date_range(queryset, desde, hasta):
    """Filtra por fecha de ingreso con límites explícitos (sin ``__date``).

    ``__date`` en MariaDB exige tener cargadas las tablas de zonas horarias;
    comparar contra límites calculados en Python funciona igual en SQLite y
    en MariaDB.
    """
    if desde:
        queryset = queryset.filter(created_at__gte=_day_start(date.fromisoformat(desde)))
    if hasta:
        try:
            limit = _day_start(date.fromisoformat(hasta) + timedelta(days=1))
        except OverflowError:
            limit = None
        if limit is not None:
            queryset = queryset.filter(created_at__lt=limit)
    return queryset


def order_attentions(queryset, orden):
    """Ordena por una clave pública (``-clave`` = descendente).

    Siempre añade ``pk`` como desempate para que la paginación sea estable.
    """
    key, descending = split_order(orden)
    if key not in SORT_FIELDS:
        key, descending = DEFAULT_SORT, False
    if key == "estado":
        queryset = queryset.annotate(
            _estado_orden=Case(
                *[When(status=value, then=Value(position))
                  for position, value in enumerate(_STATUS_PROGRESS)],
                default=Value(len(_STATUS_PROGRESS)),
                output_field=IntegerField(),
            )
        )
    fields = [f"-{name}" if descending else name for name in SORT_FIELDS[key]]
    fields.append("-pk" if descending else "pk")
    return queryset.order_by(*fields)


def attentions_listing_queryset(request, config):
    """Atenciones del listado/exportación: alcance + filtros + orden.

    No aplica la búsqueda ``q`` (la aplica quien llama con ``apply_search``,
    porque en SQLite devuelve una lista). Devuelve ``(queryset, params)``.
    """
    params = parse_list_params(request)
    queryset = filtered_queryset(request, config).select_related(
        "neighbor", "attention_type", "attention_subtype",
        "delegation", "responsible",
    )
    if params["tipo"]:
        queryset = queryset.filter(attention_type_id=params["tipo"])
    if params["subtipo"]:
        queryset = queryset.filter(attention_subtype_id=params["subtipo"])
    # La delegación solo se filtra en la vista "Todas": en cualquier otro caso
    # la delegación activa ya acota el listado (y el control no se muestra).
    if params["delegacion"] and get_active_delegation(request) == ALL_DELEGATIONS:
        queryset = queryset.filter(delegation_id=params["delegacion"])
    if params["responsable"]:
        queryset = queryset.filter(responsible_id=params["responsable"])
    queryset = _apply_date_range(queryset, params["desde"], params["hasta"])
    return order_attentions(queryset, params["orden"]), params


def list_querystring(params, **overrides):
    """``?a=1&b=2`` con los parámetros vigentes (vacíos omitidos).

    ``overrides`` reemplaza valores; pasar ``None`` quita el parámetro.
    """
    data = {key: params.get(key, "") for key in LIST_PARAM_KEYS}
    data.update(overrides)
    clean = {key: value for key, value in data.items() if value not in (None, "")}
    return f"?{urlencode(clean)}" if clean else ""


def _type_options(selected):
    options = []
    for item in AttentionType.alive.order_by("code"):
        label = item.name if item.status == "Activo" else f"{item.name} (inactivo)"
        options.append({"value": str(item.pk), "label": label,
                        "selected": str(item.pk) == selected})
    return options


def _subtype_options(selected, type_pk):
    queryset = AttentionSubtype.alive.select_related("attention_type").order_by("code")
    if type_pk:
        queryset = queryset.filter(attention_type_id=type_pk)
    options = []
    for item in queryset:
        label = item.name if item.status == "Activo" else f"{item.name} (inactivo)"
        if not type_pk:
            label = f"{item.attention_type.name} › {label}"
        options.append({"value": str(item.pk), "label": label,
                        "selected": str(item.pk) == selected})
    return options


def _responsible_options(request, config, selected):
    visible = scoped_queryset(request, config).order_by().values("responsible_id")
    users = User.objects.filter(pk__in=visible).order_by(
        "first_name", "last_name", "username",
    )
    return [
        {"value": str(user.pk), "label": user.get_full_name() or user.username,
         "selected": str(user.pk) == selected}
        for user in users
    ]


def _delegation_options(request, selected):
    return [
        {"value": str(item.pk), "label": _deleg_label(item),
         "selected": str(item.pk) == selected}
        for item in allowed_delegations(request.user)
    ]



class _AttentionFormContext:
    """Contexto extra compartido por las vistas de formulario de atención."""

    def attention_form_context(self, form):
        editing = self.mode != "create"
        if editing and self.object is not None:
            responsible_name = self.object.responsible_name
            date_text = self.object.created_at_display
            neighbor_label = (
                f"{self.object.neighbor.identifier} | {self.object.neighbor.full_name}"
                if self.object.neighbor_id else ""
            )
        else:
            responsible_name = (
                self.request.user.get_full_name() or self.request.user.username
            )
            date_text = "automáticas"
            neighbor_label = ""
        subtype_type_map = {
            subtype.pk: subtype.attention_type_id
            for subtype in form.fields["attention_subtype"].queryset
        }
        # Delegación: select (coordinador/superadmin al crear) o texto fijo.
        fixed = form.fixed_delegation
        if editing and self.object is not None:
            fixed = self.object.delegation
        fixed_label = ""
        if fixed is not None:
            suffix = ""
            if not editing and "delegation" not in form.fields:
                suffix = "  [Asignada por perfil]"
            fixed_label = f"{fixed.code} | Delegación {fixed.name}{suffix}"
        return {
            "subtype_type_map": subtype_type_map,
            "responsible_name": responsible_name,
            "attention_date_text": date_text,
            "neighbor_label": neighbor_label,
            "fixed_delegation_label": fixed_label,
            "delegation_is_select": "delegation" in form.fields,
            "status_in_form": "status" in form.fields,
        }



class AttentionListView(CrudListView):
    """Listado de atenciones: búsqueda, filtros, orden por columna y página.

    Parámetros GET (todos opcionales y validados): ``q``, ``estado``, ``tipo``,
    ``subtipo``, ``delegacion`` (solo en la vista "Todas"), ``responsable``,
    ``desde``/``hasta`` (AAAA-MM-DD), ``orden`` (clave o ``-clave``),
    ``por_pagina`` (5, 10, 25 o 50) y ``page``.
    """

    template_name = "portal/crud/attention_list.html"

    def get_context_data(self, **kwargs):
        # Se salta CrudListView.get_context_data a propósito: ese método pagina
        # con ``config.per_page`` fijo y sin los filtros nuevos, y repetir la
        # consulta solo para descartarla sería trabajo desperdiciado. Este
        # método entrega las mismas claves de contexto que usa la plantilla.
        context = super(CrudListView, self).get_context_data(**kwargs)
        config = self.config
        request = self.request

        queryset, params = attentions_listing_queryset(request, config)
        rows_all = apply_search(queryset, config.search_fields, params["q"])

        page_size = int(params["por_pagina"] or DEFAULT_PAGE_SIZE)
        paginator = Paginator(rows_all, page_size)
        page_obj = paginator.get_page(request.GET.get("page"))

        rows = [
            {
                "pk": obj.pk,
                "cells": [resolve_cell(obj, column) for column in config.columns],
            }
            for obj in page_obj.object_list
        ]

        page_links = []
        for number in paginator.get_elided_page_range(
            page_obj.number, on_each_side=2, on_ends=1,
        ):
            if number == Paginator.ELLIPSIS:
                page_links.append({"ellipsis": True})
            else:
                page_links.append({
                    "number": number,
                    "url": list_querystring(params, page=number),
                    "current": number == page_obj.number,
                })

        # Encabezados ordenables: un clic ordena ascendente y el siguiente
        # invierte; ``aria-sort`` informa a los lectores de pantalla.
        current_key, current_desc = split_order(params["orden"] or DEFAULT_SORT)
        headers = []
        for column in config.columns:
            key = COLUMN_SORT_KEY.get(column.field_name)
            is_current = key == current_key
            next_order = f"-{key}" if is_current and not current_desc else key
            headers.append({
                "label": column.header,
                "url": list_querystring(params, orden=next_order, page=None) if key else "",
                "sortable": bool(key),
                "current": is_current,
                "direction": ("desc" if current_desc else "asc") if is_current else "",
                "aria_sort": (
                    "descending" if current_desc else "ascending"
                ) if is_current else "none",
            })

        active = get_active_delegation(request)
        role = get_role(request.user)
        filters_active = any(
            params[key] for key in (
                "q", "estado", "tipo", "subtipo", "delegacion",
                "responsable", "desde", "hasta",
            )
        )
        base_search = reverse(config.module_key)

        context.update(
            {
                "active_module": config.module_key,
                "config": config,
                "rows": rows,
                "columns": config.columns,
                "headers": headers,
                "page_obj": page_obj,
                "page_links": page_links,
                "prev_url": list_querystring(
                    params, page=page_obj.previous_page_number() if page_obj.has_previous() else None,
                ),
                "next_url": list_querystring(
                    params, page=page_obj.next_page_number() if page_obj.has_next() else None,
                ),
                "showing_from": page_obj.start_index(),
                "showing_to": page_obj.end_index(),
                "total": paginator.count,
                "query": params["q"],
                "estado": params["estado"],
                "status_options": config.status_options,
                "empty_kind": "results" if filters_active else "records",
                "search_url": base_search,
                "export_url": reverse(f"{config.module_key}_export")
                + list_querystring(params, page=None, por_pagina=None),
                "new_url": reverse(f"{config.module_key}_new"),
                "detail_url_name": f"{config.module_key}_detail",
                "edit_url_name": f"{config.module_key}_edit",
                "delete_url_name": f"{config.module_key}_delete",
                # --- filtros avanzados ---
                "type_options": _type_options(params["tipo"]),
                "subtype_options": _subtype_options(params["subtipo"], params["tipo"]),
                "show_delegation_filter": active == ALL_DELEGATIONS,
                "delegation_options": (
                    _delegation_options(request, params["delegacion"])
                    if active == ALL_DELEGATIONS else []
                ),
                "show_responsible_filter": bool(role) and role.scope != "own_records",
                "responsible_options": (
                    _responsible_options(request, config, params["responsable"])
                    if role and role.scope != "own_records" else []
                ),
                "date_from": params["desde"],
                "date_to": params["hasta"],
                "date_error": params["date_error"],
                "current_order": params["orden"],
                "page_size": page_size,
                "page_size_options": PAGE_SIZE_OPTIONS,
                "scope_note": scope_note(request),
                **self.access_flags(),
            }
        )
        return context


class AttentionCreateView(_AttentionFormContext, CrudCreateView):
    """Ingreso de atención: responsable = usuario de sesión, N° automático."""

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.attention_form_context(context["form"]))
        return context

    def form_valid(self, form):
        obj = form.save(commit=False)
        obj.responsible = self.request.user
        obj.save()
        audit(
            self.request, "create", module="attentions",
            object_repr=str(obj), detail=f"Creado por {self.request.user}",
        )
        messages.success(
            self.request,
            f"Atención {obj.number} ingresada correctamente en "
            f"{obj.delegation.code} | Delegación {obj.delegation.name}. "
            "La fecha y la hora se asignaron automáticamente.",
        )
        return redirect(self.get_success_url())


class AttentionUpdateView(_AttentionFormContext, CrudUpdateView):
    """Edición: solo tipo, sub atención, descripción y estado."""

    template_name = "portal/crud/attention_form.html"

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.attention_form_context(context["form"]))
        return context

    def form_valid(self, form):
        obj = form.save()
        audit(
            self.request, "edit", module="attentions",
            object_repr=str(obj), detail=f"Editado por {self.request.user}",
        )
        messages.success(
            self.request,
            f"Cambios guardados. La atención {obj.number} se actualizó y el "
            "cambio quedó en la auditoría.",
        )
        return redirect(self.get_success_url())


class AttentionDetailView(_AttentionFormContext, CrudDetailView):
    """Consulta de una atención (solo lectura)."""

    def get(self, request, *args, **kwargs):
        self.object = self.get_object()
        form = self.config.form_class(instance=self.object, request=request)
        for field in form.fields.values():
            field.widget.attrs["disabled"] = True
        return self.render_to_response(self.get_context_data(form=form))

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context.update(self.attention_form_context(context["form"]))
        return context



class AttentionExportView(LoginRequiredMixin, ModulePermissionMixin, CrudBaseMixin, View):
    """Exporta a Excel las atenciones filtradas y con alcance."""

    action = "export"

    def get(self, request, *args, **kwargs):
        # Mismos filtros y orden que el listado; se exportan TODAS las filas
        # del resultado (sin paginar).
        queryset, params = attentions_listing_queryset(request, self.config)
        rows_all = list(apply_search(queryset, self.config.search_fields, params["q"]))

        headers = ["N°", "Fecha y hora", "Solicitante", "Tipo", "Sub atención",
                   "Delegación", "Responsable", "Estado"]

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Atenciones"

        for index, header in enumerate(headers, start=1):
            cell = sheet.cell(row=1, column=index, value=header)
            cell.fill = HEADER_FILL
            cell.font = HEADER_FONT
            cell.alignment = HEADER_ALIGN
            cell.border = BORDER

        for row_index, obj in enumerate(rows_all, start=2):
            values = [
                obj.number,
                obj.created_at_display,
                obj.solicitante_name,
                obj.type_name,
                obj.subtype_name,
                obj.delegation_label,
                obj.responsible_name,
                obj.status_display,
            ]
            for col_index, value in enumerate(values, start=1):
                cell = sheet.cell(row=row_index, column=col_index, value=value)
                cell.border = BORDER

        for index, header in enumerate(headers, start=1):
            letter = get_column_letter(index)
            sheet.column_dimensions[letter].width = max(len(header) + 2, 14)

        response = HttpResponse(
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        response["Content-Disposition"] = 'attachment; filename="atenciones.xlsx"'
        workbook.save(response)
        return response


class AttentionNeighborsView(LoginRequiredMixin, ModulePermissionMixin, View):
    """Búsqueda JSON de solicitantes (máx. 10) para el ingreso de atención."""

    module = "attentions"
    action = "view"

    def get(self, request):
        query = (request.GET.get("q") or "").strip()
        results = []
        if query:
            neighbors = (
                Neighbor.alive.filter(full_name__icontains=query)
                | Neighbor.alive.filter(identifier__icontains=query)
            )
            for neighbor in neighbors.order_by("full_name")[:10]:
                results.append({
                    "pk": neighbor.pk,
                    "label": f"{neighbor.identifier} | {neighbor.full_name}",
                })
        return JsonResponse({"results": results})
